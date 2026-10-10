"""Build the Word edition of the client review documents (docs/client_review/0*.md -> .docx).

    python scripts/build_review_pack.py [--out DIR]

Renders the SVG diagrams to PNG with the installed Chrome / Edge (scripts/ui_tour/render_svg.js, needs Node and the
browser tour's node_modules), then writes one independent .docx per document - each stands on its own, with no
reference to the others. The Markdown files stay the source of truth; the .docx files are generated (git-ignored).
"""

from __future__ import annotations

import argparse
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "client_review"
DOCS = ["01_SOLUTION_ARCHITECTURE.md", "02_SECURITY_TOOL_INTEGRATIONS.md",
        "03_DATA_FLOW_AND_PROTECTION.md", "04_AI_LLM_SECURITY.md", "05_SECURITY_RISK_AND_GOVERNANCE.md",
        "06_DEPLOYMENT_AND_VALIDATION.md"]
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*\s][^*]*\*)")


def render_diagrams(out: Path) -> dict[str, Path]:
    svgs = sorted((SRC / "diagrams").glob("*.svg"))
    png_dir = out / "diagrams"
    script = ROOT / "scripts" / "ui_tour" / "render_svg.js"
    subprocess.run(["node", str(script), str(png_dir), *map(str, svgs)], check=True, capture_output=True, text=True)
    return {s.name: png_dir / s.name.replace(".svg", ".png") for s in svgs}


def add_runs(par, text: str) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        elif part.startswith("[") and "](" in part:
            label = part[1:part.index("](")]
            par.add_run(label).underline = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def shade(cell, hex_fill: str) -> None:
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc.append(shd)


def add_table(doc, rows: list[list[str]]) -> None:
    header, body = rows[0], [r for r in rows[1:] if not all(re.fullmatch(r":?-{3,}:?", c.strip()) for c in r)]
    cols = len(header)
    blank_header = not any(h.strip() for h in header)       # "| | |" header: a key / value table
    t = doc.add_table(rows=0 if blank_header else 1, cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate([] if blank_header else header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        add_runs(cell.paragraphs[0], h.strip())
        for r in cell.paragraphs[0].runs:
            r.bold = True
        shade(cell, "E0ECFF")
    for row in body:
        cells = t.add_row().cells
        for i in range(cols):
            cells[i].text = ""
            add_runs(cells[i].paragraphs[0], (row[i] if i < len(row) else "").strip())
    # widths follow content: a short "#" or "Status" column should not take as much room as "Controls in place"
    all_rows = ([] if blank_header else [header]) + body
    weight = [max(7, min(60, max((len(r[c].strip()) if c < len(r) else 0) for r in all_rows))) for c in range(cols)]
    total = sum(weight)
    t.autofit = False
    for row in t.rows:
        for c, cell in enumerate(row.cells):
            cell.width = Cm(17.0 * weight[c] / total)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)
    doc.add_paragraph()


def convert(md: Path, doc, pngs: dict[str, Path]) -> None:
    lines = md.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(block))
            r.font.name = "Consolas"
            r.font.size = Pt(8)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, rows)
            continue
        elif m := re.match(r"!\[([^\]]*)\]\(diagrams/([^)]+)\)", line.strip()):
            png = pngs.get(m.group(2))
            if png and png.exists():
                doc.add_picture(str(png), width=Cm(16.5))
                cap = doc.add_paragraph(m.group(1))
                cap.runs[0].italic = True
        elif m := re.match(r"^(#{1,4})\s+(.*)", line):
            doc.add_heading(m.group(2), level=min(len(m.group(1)) - 1, 3) if len(m.group(1)) > 1 else 0)
        elif re.match(r"^\s*[-*]\s+", line) or re.match(r"^\s*\d+\.\s+", line):
            numbered = bool(re.match(r"^\s*\d+\.\s+", line))
            text = re.sub(r"^\s*(?:[-*]|\d+\.)\s+", "", line)
            text = re.sub(r"^\[ \]\s*", "☐  ", text)          # checklist box
            while i + 1 < len(lines) and re.match(r"^\s{2,}\S", lines[i + 1])                     and not re.match(r"^\s*(?:[-*]|\d+\.)\s+", lines[i + 1]):
                i += 1
                text += " " + lines[i].strip()           # a wrapped list item continues on indented lines
            add_runs(doc.add_paragraph(style="List Number" if numbered else "List Bullet"), text)
        elif line.strip():
            par = [line.strip()]
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"^(\||#|```|!\[|\s*[-*]\s|\s*\d+\.\s)", lines[i + 1]):
                i += 1
                par.append(lines[i].strip())
            add_runs(doc.add_paragraph(), " ".join(par))
        i += 1


def new_document():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(2)
    for lvl, size in ((0, 20), (1, 14), (2, 12), (3, 11)):
        h = doc.styles["Title" if lvl == 0 else f"Heading {lvl}"]
        h.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        h.font.size = Pt(size)
    return doc


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--out", default=str(SRC / "word"))
    args = ap.parse_args()
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    pngs = render_diagrams(out)
    for stale in ("00_REVIEW_PACK_INDEX.docx", "Agentic_SOC_Review_Pack.docx"):   # earlier builds; documents stand alone
        (out / stale).unlink(missing_ok=True)
    for name in DOCS:
        doc = new_document()
        convert(SRC / name, doc, pngs)
        target = out / name.replace(".md", ".docx")
        doc.save(target)
        print(target)
    return 0


if __name__ == "__main__":
    sys.exit(main())
