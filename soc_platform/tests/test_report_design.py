"""Reports are fit to send as they are: the house style (cover, tiles, charts, styled tables, running header, page X of
Y, classification) in every Word report and deck, numbers shown as words when missing - never 'None' or 'n/a' - and
chart axes a reader expects. The figures themselves are tested elsewhere; this pins the presentation."""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation

from soc_platform.connectors.registry import ConnectorRegistry
from soc_platform.reporting import design as D
from soc_platform.reporting.charts import nice_step

BAD_TEXT = ("None", "n/a", "nan", "undefined")


def test_numbers_read_the_same_everywhere_and_missing_values_are_words():
    assert D.fmt(None) == "not yet measured" and D.fmt("") == "not yet measured"
    assert D.fmt(1234) == "1,234" and D.fmt(12.34, "%") == "12.3%" and D.fmt(3.0, "days") == "3 days"
    assert D.fmt(True) == "yes" and D.fmt(None, missing="–") == "–"


@pytest.mark.parametrize("peak, step", [(3, 1), (4, 1), (6, 2), (12, 5), (86, 20), (240, 50), (1300, 500)])
def test_chart_axes_use_steps_a_reader_expects(peak, step):
    assert nice_step(peak) == step


@pytest.fixture()
def services(session, tmp_path):
    from soc_platform.domains.incident.service import IncidentService
    from soc_platform.domains.phishing.service import PhishingService
    from soc_platform.domains.vulnerability.service import VulnerabilityService

    reg = ConnectorRegistry.all_fake()
    vm = VulnerabilityService(session, reg)
    vm.refresh()
    inc = IncidentService(session, reg)
    inc.ingest()
    for case in inc.cluster():                                     # investigations create the recommended actions
        inc.investigate(case.id, narrate=False)
    return vm, inc, PhishingService(session, reg, org_domains=["acme-demo.com"], raw_dir=tmp_path / "raw")


def _docx_text(path: Path) -> str:
    d = Document(str(path))
    parts = [p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells]
    parts += [p.text for s in d.sections for p in s.header.paragraphs + s.footer.paragraphs]
    return "\n".join(parts)


def _words(text: str) -> set[str]:
    import re

    return set(re.findall(r"[A-Za-z/]+", text))


def test_word_reports_carry_the_house_style(session, services, tmp_path):
    from soc_platform.reporting.reports import ReportService

    vm, _, _ = services
    rs = ReportService(session, tmp_path / "out")
    for run in (rs.daily_exposure(vm, by="t"), rs.weekly_vm(vm, by="t")):
        path = Path(run.path)
        text = _docx_text(path)
        assert D.CLASSIFICATION in text and "page" in text                          # running header and footer
        xml = zipfile.ZipFile(path).read("word/footer1.xml").decode()
        assert "PAGE" in xml and "NUMPAGES" in xml                                  # page X of Y fields
        media = [n for n in zipfile.ZipFile(path).namelist() if n.startswith("word/media/")]
        assert len(media) >= 2                                                      # charts, not only tables
        assert not (_words(text) & set(BAD_TEXT)), (run.kind, _words(text) & set(BAD_TEXT))
        doc = Document(str(path))
        header_fills = {c._tc.tcPr.xpath("./w:shd/@w:fill")[0] for t in doc.tables[2:] for c in t.rows[0].cells
                        if c._tc.tcPr is not None and c._tc.tcPr.xpath("./w:shd/@w:fill")}
        assert D.NAVY in header_fills                                               # styled table headers
    assert "Act today" in _docx_text(Path(rs.daily_exposure(vm, by="t").path))


def test_decks_are_widescreen_with_native_charts_and_a_footer_on_every_slide(session, services, tmp_path):
    from soc_platform.reporting.reports import ReportService

    vm, inc, ph = services
    run = ReportService(session, tmp_path / "out").weekly_management_deck(vm, inc, ph, by="t")
    prs = Presentation(run.path)
    assert round(prs.slide_width / prs.slide_height, 2) == round(16 / 9, 2)
    charts = sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart)
    assert charts >= 3
    for i, slide in enumerate(prs.slides):
        text = " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
        text += " ".join(c.text for sh in slide.shapes if sh.has_table for r in sh.table.rows for c in r.cells)
        assert not (_words(text) & set(BAD_TEXT)), (i, _words(text) & set(BAD_TEXT))
        if i:
            assert D.CLASSIFICATION in text                                         # footer on content slides


def test_built_reports_use_the_house_style_too(session, services, tmp_path):
    from soc_platform.reporting import builder as B

    reg = ConnectorRegistry.all_fake()
    deck = B.build_report(session, reg, B.get_template(session, "board_monthly"), tmp_path / "b", llm=None, by="t")
    doc = B.build_report(session, reg, B.get_template(session, "ciso_weekly"), tmp_path / "c", llm=None, by="t")
    from soc_platform.domains.vulnerability.models import ReportRun

    prs = Presentation(session.get(ReportRun, deck["id"]).path)
    assert round(prs.slide_width / prs.slide_height, 2) == round(16 / 9, 2)
    assert sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart) >= 2
    text = _docx_text(Path(session.get(ReportRun, doc["id"]).path))
    assert D.CLASSIFICATION in text and "Evidence" in text


def test_the_approvals_summary_counts_by_type_within_the_callers_scope(session, services):
    from soc_platform.api.app import actions_summary
    from soc_platform.core.auth import Principal, Role

    lead = Principal("lena", "lena", frozenset({Role.LEAD}))
    vm_only = Principal("vic", "vic", frozenset({Role.ANALYST}), domains=frozenset({"vulnerability"}))
    every = actions_summary(status="recommended,pending_approval", p=lead, s=session)
    assert every["total"] and sum(every["by_type"].values()) == every["total"]
    assert list(every["by_type"].values()) == sorted(every["by_type"].values(), reverse=True)
    scoped = actions_summary(status="recommended,pending_approval", p=vm_only, s=session)
    assert sum(scoped["by_type"].values()) == scoped["total"] == scoped["by_domain"].get("vulnerability", 0)
