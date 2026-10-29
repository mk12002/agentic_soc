"""Audit trail exports (JSON Lines, JSON, CSV, Excel, Word) with filters, and the facts the audit page states."""

from __future__ import annotations

import csv
import io
import json
import os
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET  # nosec B405 - parses the platform's own output in a test

import pytest

from soc_platform.core.audit import GENESIS, AuditLog, compute_hash
from soc_platform.core.audit_export import COLUMNS, export, info
from soc_platform.core.watchlists import read_rows

ROOT = Path(__file__).resolve().parents[2]
SECRET = "test-secret-0123456789abcdef0123456789"
_ENV = ("SOC_AUTH_MODE", "SOC_DEV_JWT_SECRET", "SOC_DATABASE_URL", "SOC_CONNECTORS_CONFIG", "SOC_REPORT_OUTPUT_DIR",
        "SOC_RAW_PAYLOAD_DIR", "SOC_ENVIRONMENT")


# ---------------------------------------------------------------------------------------------- unit


def _seed(session):
    log = AuditLog(session)
    log.append(actor_type="human", actor_id="lena@example.com", event_type="case.assigned", subject_type="case",
               subject_id="=SUM(A1:A9)", payload={"note": "line\x01break", "n": 3})
    log.append(actor_type="agent", actor_id="agent:incident", event_type="case.assessed", subject_type="case",
               subject_id="c2", payload={"verdict": "true_positive"})
    log.append(actor_type="system", actor_id="system:retention", event_type="retention.run", subject_type="job",
               subject_id="retention", payload={"x": "y" * 40_000})


def test_csv_cells_are_never_run_as_formulas_and_every_record_is_one_row(session):
    _seed(session)
    body, media, name = export(session, "csv", by="aud", filters={})
    assert media == "text/csv" and name.endswith(".csv")
    rows = list(csv.reader(io.StringIO(body.decode("utf-8-sig"))))
    assert rows[0] == COLUMNS and len(rows) == 4
    assert rows[1][COLUMNS.index("subject_id")] == "'=SUM(A1:A9)"            # a spreadsheet shows it as text
    assert json.loads(rows[1][COLUMNS.index("payload")])["n"] == 3


def test_the_excel_workbook_is_well_formed_and_holds_every_record_and_its_details(session):
    _seed(session)
    body, media, name = export(session, "xlsx", by="aud", filters={})
    assert name.endswith(".xlsx") and "spreadsheetml" in media
    with zipfile.ZipFile(io.BytesIO(body)) as z:
        assert {"[Content_Types].xml", "xl/workbook.xml", "xl/styles.xml", "xl/worksheets/sheet1.xml"} <= set(z.namelist())
        for part in z.namelist():
            ET.fromstring(z.read(part))  # nosec B314 - every part is well-formed XML (a control character would break it)
    rows = read_rows(body, name)
    log = [r for w, r in rows if w.startswith("Audit log")]
    about = {r[0]: r[1] for w, r in rows if w.startswith("About this export")}
    assert log[0] == COLUMNS and len(log) == 4
    assert log[1][COLUMNS.index("subject_id")] == "=SUM(A1:A9)"              # inline text: never a formula in .xlsx
    assert "cut" in log[3][COLUMNS.index("payload")] and len(log[3][COLUMNS.index("payload")]) < 32_767
    assert about["Records in this export"] == "3" and about["Whole chain verified at export"] == "yes"


def test_an_unfiltered_export_can_be_re_verified_outside_the_platform(session):
    _seed(session)
    doc = json.loads(export(session, "json", by="aud", filters={})[0])
    prev = GENESIS
    from datetime import datetime

    for r in doc["records"]:
        assert r["prev_hash"] == prev
        assert r["hash"] == compute_hash(prev, datetime.fromisoformat(r["ts_utc"]), r["actor_type"], r["actor_id"],
                                         r["event_type"], r["subject_type"], r["subject_id"], r["payload"])
        prev = r["hash"]
    assert doc["export"]["chain_verification"]["ok"] and doc["export"]["records"] == 3
    assert doc["export"]["selection"].startswith("the whole log")


def test_filters_select_by_event_type_prefix_actor_and_time_window(session):
    _seed(session)
    pick = lambda **f: json.loads(export(session, "json", by="aud", filters=f)[0])
    assert [r["event_type"] for r in pick(event_type="case.*")["records"]] == ["case.assigned", "case.assessed"]
    assert pick(event_type="case.*")["export"]["selection"].startswith("filtered")
    assert [r["actor_id"] for r in pick(actor_id="agent:incident")["records"]] == ["agent:incident"]
    assert pick(event_type="case_*")["records"] == []                            # "_" is literal, not a wildcard
    assert len(pick(start="2000-01-01", end="2999-12-31")["records"]) == 3
    assert pick(end="2000-01-01")["records"] == []
    with pytest.raises(ValueError, match="not a date"):
        pick(start="last week")
    with pytest.raises(ValueError, match="before its start"):
        pick(start="2026-02-01", end="2026-01-01")
    with pytest.raises(ValueError, match="unknown format"):
        export(session, "pdf", by="aud", filters={})


def test_the_word_report_opens_and_summarises_the_selection(session):
    from docx import Document

    _seed(session)
    body, _, name = export(session, "docx", by="aud", filters={})
    d = Document(io.BytesIO(body))
    text = "\n".join(p.text for p in d.paragraphs) + " ".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    assert name.endswith(".docx") and "Audit trail" in text and "case.assessed" in text and "verified" in text


def test_the_audit_facts_are_read_from_the_log_and_the_settings(session):
    from soc_platform.config import get_settings

    _seed(session)
    s = get_settings()
    out = info(session, s)
    assert out["records"] == 3 and out["by_actor_type"] == {"agent": 1, "human": 1, "system": 1}
    assert {c["category"] for c in out["by_category"]} == {"case", "retention"}
    days = {r["data"]: r["days"] for r in out["retention"]["other"]}
    assert days["Model prompts and answers"] == s.llm_log_retention_days
    assert days["Access log (every request)"] == s.access_log_retention_days
    assert "auditor" in out["access"]["export"] and "analyst" not in out["access"]["export"]
    assert [f["format"] for f in out["formats"]] == ["jsonl", "json", "csv", "xlsx", "docx"]


# ---------------------------------------------------------------------------------------------- API


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    from soc_platform.config import get_settings
    from soc_platform.core import db as dbm

    saved = {k: os.environ.get(k) for k in _ENV}
    tmp = tmp_path_factory.mktemp("ax")
    os.environ.update({"SOC_AUTH_MODE": "dev", "SOC_DEV_JWT_SECRET": SECRET, "SOC_ENVIRONMENT": "test",
                       "SOC_DATABASE_URL": f"sqlite:///{tmp / 'ax.db'}",
                       "SOC_CONNECTORS_CONFIG": str(ROOT / "config" / "connectors.yaml"),
                       "SOC_REPORT_OUTPUT_DIR": str(tmp / "reports"), "SOC_RAW_PAYLOAD_DIR": str(tmp / "raw")})
    get_settings.cache_clear()
    dbm._default = None
    from fastapi.testclient import TestClient

    from soc_platform.api import app as appmod

    appmod.registry.cache_clear()
    try:
        c = TestClient(appmod.app)
        assert c.post("/api/v1/incidents/run", headers=tok(c, "seed", "analyst")).status_code == 200
        yield c
    finally:
        appmod.registry.cache_clear()
        for k, v in saved.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v
        get_settings.cache_clear()
        dbm._default = None


def tok(client, user, roles, domains=""):
    t = client.get(f"/api/v1/dev/token?user={user}&roles={roles}&domains={domains}").json()["token"]
    return {"Authorization": f"Bearer {t}"}


def test_every_format_downloads_and_each_export_is_itself_audited(client):
    aud = tok(client, "audrey", "auditor")
    for fmt in ("jsonl", "json", "csv", "xlsx", "docx"):
        r = client.get(f"/api/v1/audit/export?format={fmt}", headers=aud)
        assert r.status_code == 200 and f'.{fmt}"' in r.headers["content-disposition"], fmt
    jsonl = client.get("/api/v1/audit/export", headers=aud).text.strip().splitlines()      # the default stays JSON Lines
    assert json.loads(jsonl[-1])["_verification"]["ok"]
    r = client.get("/api/v1/audit/export?format=csv&event_type=audit.*", headers=aud)
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    formats = [json.loads(x["payload"]).get("format") for x in rows if x["event_type"] == "audit.export"]
    assert formats[:5] == ["jsonl", "json", "csv", "xlsx", "docx"] and all(x["actor_id"] == "audrey" for x in rows)


def test_exports_need_the_export_permission_and_all_domain_scope(client):
    assert client.get("/api/v1/audit/export?format=csv", headers=tok(client, "ana", "analyst")).status_code == 403
    assert client.get("/api/v1/audit/export?format=csv",
                      headers=tok(client, "vic", "lead", "incident")).status_code == 403
    aud = tok(client, "audrey", "auditor")
    assert client.get("/api/v1/audit/export?format=pdf", headers=aud).status_code == 422
    assert client.get("/api/v1/audit/export?format=csv&start=yesterday", headers=aud).status_code == 400
    info_ = client.get("/api/v1/audit/info", headers=tok(client, "ana", "analyst")).json()     # readers see the facts
    assert info_["records"] == client.get("/api/v1/audit/verify", headers=aud).json()["records"] > 0
