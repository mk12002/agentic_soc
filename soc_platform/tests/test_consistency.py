"""Whole-system consistency: the same figure is the same everywhere, re-running changes nothing, the LLM never
changes a figure, no route crashes or leaks, and every stored reference resolves.

These tests look at the platform from several independent angles at once, so a number computed one way on a
dashboard and another way in a report, brief, answer or document cannot drift apart unnoticed.
"""

from __future__ import annotations

import io
import json
import os
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import pytest

from soc_platform.tests.conftest import ESTATES, estate_env

ROOT = Path(__file__).resolve().parents[2]


# --------------------------------------------------------------------------------------------- HTTP estate
@pytest.fixture(scope="module", params=ESTATES)
def api(request, tmp_path_factory, estate_configs):
    """One estate (built-in or seeded variant) loaded over HTTP exactly as the console does it."""
    cfg = estate_configs[request.param]
    tmp = tmp_path_factory.mktemp(f"consistency-{cfg['name']}")
    env = {"SOC_AUTH_MODE": "dev", "SOC_DEV_JWT_SECRET": "c" * 40, "SOC_ENVIRONMENT": "test",
           "SOC_DATABASE_URL": f"sqlite:///{tmp / 'c.db'}", "SOC_REPORT_OUTPUT_DIR": str(tmp / "rep"),
           "SOC_RAW_PAYLOAD_DIR": str(tmp / "raw"), "SOC_LLM_PROVIDER": "none",
           "SOC_RATE_LIMIT_RPS": "100000", "SOC_RATE_LIMIT_BURST": "100000"}
    saved = {k: os.environ.get(k) for k in env}
    os.environ.update(env)
    scope = estate_env(cfg)
    scope.__enter__()
    from soc_platform.config import get_settings
    from soc_platform.core import db as dbm

    get_settings.cache_clear()
    dbm._default = None
    from fastapi.testclient import TestClient

    from soc_platform.api import app as appmod

    appmod.registry.cache_clear()
    limiter = appmod._limiter                      # created at first import: replace for this module, restore after
    appmod._limiter = appmod._RateLimiter(1e9, 10**9)
    c = TestClient(appmod.app, raise_server_exceptions=False)

    def H(user, roles, domains=""):
        return {"Authorization": "Bearer " + c.get(f"/api/v1/dev/token?user={user}&roles={roles}&domains={domains}").json()["token"]}

    lead, ops = H(cfg["lead"], "lead"), H(f"ops@{cfg['org']}", "automation_admin")
    for step in ("/api/v1/vm/refresh", "/api/v1/incidents/run", "/api/v1/phishing/ingest"):
        assert c.post(step, headers=lead).status_code == 200, step
    for f in cfg["uploads"]:
        r = c.post("/api/v1/phishing/submit", headers=lead, files={"file": (f, (Path(cfg["corpus_dir"]) / f).read_bytes())})
        assert r.status_code == 200, f
    assert c.post("/api/v1/vm/campaigns", headers=lead, json={"cve": cfg["campaign_cve"], "notify_via": "ticket"}).status_code == 200
    assert c.post("/api/v1/vm/misconfigurations/route", headers=lead).status_code == 200
    assert c.post("/api/v1/intelligence/refresh", headers=lead).status_code == 200
    for j in ("intelligence", "follow_up", "retention"):
        assert c.post(f"/api/v1/jobs/{j}/run", headers=ops).status_code == 200
    yield {"c": c, "H": H, "lead": lead, "app": appmod.app, "cfg": cfg}
    appmod._limiter = limiter
    scope.__exit__(None, None, None)
    for k, v in saved.items():
        if v is None:
            os.environ.pop(k, None)
        else:
            os.environ[k] = v
    get_settings.cache_clear()
    dbm._default = None


def J(r):
    assert r.status_code == 200, (str(r.request.url), r.status_code, r.text[:300])
    return r.json()


def _aware(iso: str) -> datetime:
    d = datetime.fromisoformat(iso)
    return d if d.tzinfo else d.replace(tzinfo=timezone.utc)


# --------------------------------------------------------------------------------------------- 1. same figure everywhere
def test_every_figure_agrees_across_every_surface(api):
    c, lead = api["c"], api["lead"]
    aud = api["H"](f"audrey@{api['cfg']['org']}", "auditor")
    ov = J(c.get("/api/v1/dashboard/overview", headers=lead))
    cases = J(c.get("/api/v1/cases", headers=lead))
    csum = J(c.get("/api/v1/cases/summary", headers=lead))
    acts = J(c.get("/api/v1/actions?status=recommended,pending_approval", headers=lead))
    asum = J(c.get("/api/v1/actions/summary?status=recommended,pending_approval", headers=lead))
    ins = J(c.get("/api/v1/intelligence/insights", headers=lead))
    brief = J(c.get("/api/v1/intelligence/brief", headers=lead))["facts"]
    vmm = J(c.get("/api/v1/vm/metrics", headers=lead))
    finds = J(c.get("/api/v1/vm/findings", headers=lead))
    mis = J(c.get("/api/v1/vm/misconfigurations", headers=lead))
    phm = J(c.get("/api/v1/phishing/metrics", headers=lead))
    sup = J(c.get("/api/v1/phishing/suppliers", headers=lead))
    hand = J(c.get("/api/v1/incidents/handover?hours=720", headers=lead))
    cov = J(c.get("/api/v1/dashboard/attack-coverage", headers=lead))
    shadow = J(c.get("/api/v1/dashboard/shadow-it", headers=lead))
    conns = J(c.get("/api/v1/dashboard/connectors", headers=lead))

    open_cases = [x for x in cases if x["status"] != "closed"]
    open_ins = [i for i in ins if i["status"] in ("new", "acknowledged")]
    open_f = [f for f in finds if f["status"] in ("open", "reopened")]
    now = datetime.now(timezone.utc)
    agree = {
        "open cases": [ov["cases"]["open"], len(open_cases), brief["open_cases_total"]],
        "open cases by domain": [ov["cases"]["open_by_domain"], dict(Counter(x["domain"] for x in open_cases))],
        "cases total": [csum["total"], len(cases)],
        "cases by domain": [csum["by_domain"], dict(Counter(x["domain"] for x in cases))],
        "awaiting approval": [ov["actions"]["pending_approval"], len(acts), asum["total"], brief["pending_approvals"]],
        "approvals by domain": [asum["by_domain"], dict(Counter(a["domain"] for a in acts))],
        "open insights": [ov["insights"]["open"], len(open_ins), brief["open_insights_total"]],
        "insights by severity": [ov["insights"]["by_severity"], dict(Counter(i["severity"] for i in open_ins))],
        "open vulnerabilities": [ov["vulnerability"]["open_findings"], vmm["open"], len(open_f), brief["vulnerability"]["open"]],
        "priority mix": [ov["vulnerability"]["by_band"], vmm["by_priority"], dict(Counter(f["priority"] for f in open_f))],
        "KEV open": [vmm["kev_open"], sum(1 for f in open_f if (f["factors"] or {}).get("kev")), brief["vulnerability"]["kev_open"]],
        "past SLA": [ov["vulnerability"]["sla_breached"], vmm["sla_breached"],
                     sum(1 for f in open_f if f["sla_due"] and _aware(f["sla_due"]) < now)],
        "internet exposed": [ov["vulnerability"]["internet_exposed"], vmm["internet_exposed_open"],
                             sum(1 for f in open_f if f["internet_exposed"])],
        "open misconfigurations": [ov["vulnerability"]["open_misconfigurations"], mis["metrics"]["open"],
                                   sum(1 for m in mis["items"] if m["status"] in ("open", "routed", "reopened", "pending_validation"))],
        "phishing reports": [phm["reported"], sum(1 for x in cases if x["domain"] == "phishing")],
        "phishing verdicts": [phm["verdict_mix"], dict(Counter(x["verdict"] for x in cases if x["domain"] == "phishing"))],
        "auto-closed": [phm["auto_closed"], sum(1 for x in cases if x["domain"] == "phishing" and x["status"] in ("closed", "awaiting_qa"))],
        "sampled for QA": [phm["sampled_for_qa"], sum(1 for x in cases if x["status"] == "awaiting_qa")],
        "supplier findings": [len(sup["findings"]), sum(1 for i in ins if i["rule"] == "supplier_risk")],
        "open incidents": [hand["open_total"], sum(1 for x in open_cases if x["domain"] == "incident")],
    }
    bad = {k: v for k, v in agree.items() if len({json.dumps(x, sort_keys=True) for x in v}) != 1}
    assert not bad, bad

    # risk: one score per entity on every surface (top list, entity risk, 360, insight titles, report facts)
    risk = J(c.get("/api/v1/intelligence/risk/top?limit=20", headers=lead))
    rep = J(c.post("/api/v1/reports/build", headers=lead, json={"spec": {"title": "r", "sections": [{"source": "risk"}]}}))
    rep_scores = {k: int(v.split("risk ")[1].split("/")[0]) for k, v in rep["sections"][0]["facts"]}
    for r in risk[:8]:
        scores = {r["score"], J(c.get(f"/api/v1/intelligence/entities/{r['entity_id']}/risk", headers=lead))["score"],
                  J(c.get(f"/api/v1/entities/{r['entity_id']}/360", headers=lead))["risk"]["score"]}
        scores |= {int(m) for i in ins if r["name"] in i["title"] for m in re.findall(r"\((\d+)/100\)", i["title"])}
        if r["name"] in rep_scores:
            scores.add(rep_scores[r["name"]])
        assert len(scores) == 1 and isinstance(r["score"], int), (r["name"], scores)

    # attack story: summary, assessment, KPIs and plan agree with each other
    pc = next(x for x in cases if x["domain"] == "phishing" and api["cfg"]["phish_subject_token"] in x["title"])
    st = J(c.get(f"/api/v1/cases/{pc['id']}/story", headers=lead))
    pending = sum(1 + len(a.get("duplicate_ids", [])) for p in st["response_plan"] for a in p["actions"] if a["approvable"])
    assert f"({pending} action(s) awaiting approval)" in st["summary"]
    reached = {s["stage"] for s in st["steps"] if s["outcome"] != "blocked"}
    assert st["assessment"]["reason"].startswith(f"{len(reached)} kill-chain stage(s) reached")
    assert sorted(st["tools"]) == sorted({t for s in st["steps"] for t in s["tools"]})
    assert st["blast_radius"]["stats"]["users_received"] == len(st["blast_radius"]["recipients"])

    # a case lists the same actions on its page and through the actions API
    cd = J(c.get(f"/api/v1/cases/{pc['id']}", headers=lead))
    assert sorted(a["id"] for a in cd["actions"]) == sorted(a["id"] for a in J(c.get(f"/api/v1/actions?case_id={pc['id']}", headers=lead)))

    # reports carry the same figures as the dashboards (builder facts and the generated documents)
    from docx import Document

    b = J(c.post("/api/v1/reports/build", headers=lead, json={"template_id": "ciso_weekly"}))
    f = {s["source"]: dict(s["facts"]) for s in b["sections"]}
    assert f["overview"]["Open cases"] == str(ov["cases"]["open"])
    assert f["overview"]["Actions awaiting approval"] == str(ov["actions"]["pending_approval"])
    assert f["overview"]["Open correlated insights"] == str(ov["insights"]["open"]) == f["correlated_findings"]["Open correlated findings"]
    assert f["detection_coverage"]["Weighted ATT&CK coverage"] == f"{cov['summary']['weighted_coverage_pct']}%"
    assert f["integrations"]["Enabled connectors"] == str(sum(1 for x in conns if x["enabled"]))
    vmr = J(c.post("/api/v1/reports/build", headers=lead, json={"template_id": "vm_weekly"}))
    assert dict(vmr["sections"][0]["facts"])["Open findings"] == str(vmm["open"])
    doc = Document(io.BytesIO(c.get(f"/api/v1/reports/{vmr['id']}/download", headers=lead).content))
    assert f"Open findings: {vmm['open']}" in "\n".join(p.text for p in doc.paragraphs)
    dex = J(c.post("/api/v1/reports/daily_exposure", headers=lead))
    ddoc = Document(io.BytesIO(c.get(f"/api/v1/reports/{dex['id']}/download", headers=lead).content))
    cells = {row.cells[0].text: row.cells[1].text for t in ddoc.tables for row in t.rows if len(row.cells) >= 2}
    assert cells.get("Open findings") == str(vmm["open"]) and cells.get("KEV-listed open") == str(vmm["kev_open"])
    assert cells.get("Past SLA") == str(vmm["sla_breached"])
    sh = J(c.post("/api/v1/reports/build", headers=lead, json={"spec": {"title": "s", "sections": [{"source": "shadow_it"}]}}))
    assert dict(sh["sections"][0]["facts"])["Unsanctioned services"] == str(shadow["summary"]["unsanctioned_services"])
    comp = J(c.post("/api/v1/reports/compliance", headers=aud))
    cq = J(c.post("/api/v1/reports/build", headers=aud, json={"template_id": "compliance_quarterly"}))
    assert dict(cq["sections"][0]["facts"])["Control tests passed"] == f"{comp['summary']['passed']} of {comp['summary']['tests']}"

    # audit: the verified chain and the log are the same records
    assert J(c.get("/api/v1/audit/verify", headers=aud))["records"] == len(J(c.get("/api/v1/audit?limit=100000", headers=aud)))


# --------------------------------------------------------------------------------------------- 2. routes x roles
def test_every_get_route_as_every_role_no_errors_no_leaks_explicit_utc(api):
    c, H, app = api["c"], api["H"], api["app"]
    lead = api["lead"]
    cases = J(c.get("/api/v1/cases", headers=lead))
    rep = J(c.post("/api/v1/reports/build", headers=lead, json={"template_id": "vm_weekly"}))
    ids = {"cid": next(x["id"] for x in cases if x["domain"] == "phishing"),
           "eid": J(c.get("/api/v1/intelligence/risk/top", headers=lead))[0]["entity_id"], "rid": rep["id"], "cve": api["cfg"]["campaign_cve"]}
    by_domain = {d: {x["id"] for x in cases if x["domain"] == d} for d in ("phishing", "incident", "vulnerability")}
    org = api["cfg"]["org"]
    roles = {"lead": lead, "analyst": H(f"alice@{org}", "analyst"), "auditor": H(f"audrey@{org}", "auditor"),
             "admin": H(f"ada@{org}", "admin"), "phishing_only": H(f"pia@{org}", "analyst", "phishing"),
             "vulnerability_only": H(f"vic@{org}", "analyst", "vulnerability"), "anonymous": {}}
    extra = {"/api/v1/entities/find": f"?kind=identity&key=upn&value={api['cfg']['focus_upn']}", "/api/v1/metrics/shadow": "?domain=phishing"}
    # on a variant estate, no response may mention the built-in estate: nothing in the code is tied to its names
    toks = api["cfg"]["original_tokens"]
    leak = re.compile("|".join(r"(?<![a-z0-9])" + re.escape(t) + r"(?![a-z0-9])" for t in toks), re.I) if toks else None
    iso = re.compile(r'"(ts_utc|[a-z_]*(?:_at|_due|_seen|since|until|when|start|end|date))":\s*"(\d{4}-\d{2}-\d{2}T[\d:.]+)(Z|[+-]\d{2}:\d{2})?"')
    problems = []
    gets = sorted({r.path for r in app.routes if hasattr(r, "methods") and "GET" in r.methods and r.path.startswith("/api")}
                  - {"/api/v1/dev/token", "/api/v1/audit/export"})
    for path in gets:
        url = path
        for k, v in ids.items():
            url = url.replace("{" + k + "}", v)
        url += extra.get(path, "")
        for name, h in roles.items():
            r = c.get(url, headers=h)
            if r.status_code >= 500:
                problems.append(f"{r.status_code} {name} {url}")
            if name == "anonymous" and r.status_code != 401 and path != "/api/v1/llm/status":
                problems.append(f"anonymous got {r.status_code} at {url}")
            if r.status_code == 200 and "json" in r.headers.get("content-type", ""):
                for m in iso.finditer(r.text):
                    if not m.group(3):
                        problems.append(f"timestamp without UTC offset: {path} {m.group(1)}")
                m = leak.search(r.text) if leak else None
                if m:
                    problems.append(f"built-in estate name '{m.group(0)}' in {name} response at {path}")
                if name.endswith("_only"):
                    mine = name.split("_")[0]
                    for d, idset in by_domain.items():
                        if d != mine and any(i in r.text for i in idset):
                            problems.append(f"LEAK: {name} sees {d} data at {url}")
        if "{" in path:                                   # unknown ids answer 4xx, never 5xx
            for bogus in ("does-not-exist", "0" * 32, "%00", "a" * 300):
                u = path
                for k in ids:
                    u = u.replace("{" + k + "}", bogus)
                r = c.get(u, headers=lead)
                if r.status_code >= 500:
                    problems.append(f"{r.status_code} for unknown id at {u[:80]}")
    assert not sorted(set(problems)), sorted(set(problems))[:20]


def test_malformed_input_never_crashes_any_write_route(api):
    if api["cfg"]["name"] != "demo":
        pytest.skip("input fuzzing does not depend on the data set: run once, on the built-in estate")
    c, H, app = api["c"], api["H"], api["app"]
    last = ("/api/v1/admin/revoke-sessions", "/api/v1/auth/logout", "/api/v1/kill-switch")
    routes = sorted({(m, r.path) for r in app.routes if hasattr(r, "methods") and r.path.startswith("/api")
                     for m in r.methods if m in {"POST", "DELETE"}}, key=lambda x: (x[1] in last, x[1]))
    bodies = [None, {}, [], "not json", {"action_ids": "x", "spec": {"sections": "x"}, "cve": 42, "request": "a" * 5000}]
    bogus = ["does-not-exist", "0" * 32, "' OR 1=1 --"]
    problems = []
    for method, path in routes:
        roles = {"lead": H(api["cfg"]["lead"], "lead"), "admin": H(f"ada@{api['cfg']['org']}", "admin")}  # fresh: revoke-sessions
        params = re.findall(r"{(\w+)}", path)
        for idv in (bogus if params else [""]):
            url = path
            for p in params:
                url = url.replace("{" + p + "}", {"verb": "approve", "kind": "bogus", "name": "bogus"}.get(p, idv))
            for body in bodies:
                for rn, h in roles.items():
                    kw = {"headers": h}
                    if isinstance(body, str):
                        kw = {"headers": {**h, "content-type": "application/json"}, "content": body.encode()}
                    elif body is not None:
                        kw["json"] = body
                    r = c.request(method, url, **kw)
                    if r.status_code >= 500:
                        problems.append(f"{r.status_code} {rn} {method} {url}")
    assert not sorted(set(problems)), sorted(set(problems))


# --------------------------------------------------------------------------------------------- 3. re-running changes nothing
def _estate(s, reg, cfg, llm=None):
    from soc_platform.core.auth import Principal, Role
    from soc_platform.domains.incident.service import IncidentService
    from soc_platform.domains.phishing.service import PhishingService
    from soc_platform.domains.vulnerability.misconfig import MisconfigurationService
    from soc_platform.domains.vulnerability.service import VulnerabilityService
    from soc_platform.intelligence.analyst import IntelligenceService

    lead = Principal(cfg["lead"], "Lead", frozenset({Role.LEAD}))
    VulnerabilityService(s, reg, llm=llm).refresh()
    mis = MisconfigurationService(s, reg)
    mis.refresh()
    inc = IncidentService(s, reg, llm=llm)
    inc.ingest()
    for c in inc.cluster():
        inc.investigate(c.id)
    ph = PhishingService(s, reg, org_domains=[cfg["org"]], llm=llm)
    for sub in ph.ingest_reported():
        ph.process(sub.id)
    for f in cfg["uploads"]:
        sub = ph.submit_raw((Path(cfg["corpus_dir"]) / f).read_bytes(), source="upload", reporter=cfg["lead"])
        ph.process(sub.id)
    VulnerabilityService(s, reg, llm=llm).create_campaign(cfg["campaign_cve"], lead, notify_via="ticket")
    mis.route(lead)
    IntelligenceService(s, llm).refresh()


def _state(s):
    from sqlalchemy import inspect, text

    from soc_platform.intelligence.risk import RiskEngine

    skip = {"audit_log", "access_log", "llm_calls", "job_runs", "enrichment_cache", "system_flags", "report_runs"}
    counts = {t: s.execute(text(f"select count(*) from {t}")).scalar() for t in inspect(s.connection()).get_table_names() if t not in skip}
    return {"counts": counts,
            "risk": {p.name: p.score for p in RiskEngine(s).top(None, 100)},
            "insights": sorted(s.execute(text("select rule || '|' || severity || '|' || title from insights")).scalars()),
            "cases": sorted(s.execute(text("select domain || '|' || title || '|' || severity || '|' || coalesce(verdict,'') || '|' || status || '|' || round(coalesce(confidence,0), 3) from cases")).scalars()),
            "actions": sorted(s.execute(text("select action_type || '|' || status || '|' || domain from action_requests")).scalars()),
            "findings": sorted(s.execute(text("select cve || '|' || asset_name || '|' || priority_band || '|' || round(priority_score, 3) from vm_findings")).scalars())}


@pytest.mark.parametrize("estate", ESTATES)
def test_rerunning_every_pipeline_and_job_changes_nothing(estate, estate_configs):
    from soc_platform.connectors.registry import ConnectorRegistry
    from soc_platform.core.db import Database
    from soc_platform.jobs import JOBS, run_job

    db = Database("sqlite://")
    db.create_all()
    reg = ConnectorRegistry.all_fake()
    snaps = []
    with estate_env(estate_configs[estate]) as cfg:
        for _ in range(2):
            with db.session() as s:
                _estate(s, reg, cfg)
            for name in JOBS:
                r = run_job(name, db=db, sleep=lambda _x: None)
                assert r is None or r.status == "ok", (name, r.error)
            with db.session() as s:
                snaps.append(_state(s))
    diffs = {k: (snaps[0][k], snaps[1][k]) for k in snaps[0] if snaps[0][k] != snaps[1][k]}
    assert not diffs, {k: v if not isinstance(v[0], dict) else {n: (v[0].get(n), v[1].get(n)) for n in set(v[0]) | set(v[1])
                                                                if v[0].get(n) != v[1].get(n)} for k, v in diffs.items()}


# --------------------------------------------------------------------------------------------- 4. LLM never changes a figure
@pytest.mark.parametrize("estate", ESTATES)
def test_llm_on_or_off_gives_identical_figures_verdicts_and_actions(estate, estate_configs):
    from soc_platform.config import Settings
    from soc_platform.connectors.registry import ConnectorRegistry
    from soc_platform.core.db import Database
    from soc_platform.llm.gateway import Completion, LLMGateway, Provider

    class Talkative(Provider):
        """Writes confident narrative citing the first evidence id - and tries to state its own figures."""

        name = "scripted"

        def complete(self, system, user, *, tier):
            ids = re.findall(r"^\[([A-Z]\d+)\]", user, re.M) or ["E1"]
            if '"calls"' in user:
                return Completion(json.dumps({"calls": [{"tool": "list_insights", "args": {}}]}), 10, 10, "m")
            return Completion(json.dumps({"summary": "Risk is 12/100 across 99 tools.", "claims": [
                {"text": "The situation is serious.", "kind": "inference", "evidence_ids": ids[:1]},
                {"text": "There are 999 affected hosts.", "kind": "fact", "evidence_ids": ids[:1]}]}), 50, 20, "m")

    reg = ConnectorRegistry.all_fake()
    results = []
    for with_llm in (False, True):
        db = Database("sqlite://")
        db.create_all()
        with estate_env(estate_configs[estate]) as cfg, db.session() as s:
            llm = LLMGateway(s, Settings(org_domains=[cfg["org"]]), provider=Talkative()) if with_llm else None
            _estate(s, reg, cfg, llm)
            st = _state(s)
            from soc_platform.core.models import Case
            texts = [c.summary or "" for c in s.query(Case)] + [cl["text"] for c in s.query(Case) for cl in (c.assessment or {}).get("claims", [])]
            results.append((st, texts))
    (off, _), (on, on_texts) = results
    diffs = {k: (off[k], on[k]) for k in off if k != "counts" and off[k] != on[k]}
    assert not diffs, diffs
    assert {k: v for k, v in off["counts"].items()} == on["counts"]
    assert not any("999" in t or "99 tools" in t for t in on_texts)          # unsupported figures never shown


# --------------------------------------------------------------------------------------------- 5. integrity invariants
def test_every_stored_reference_resolves(api):
    from sqlalchemy import select

    from soc_platform.core import db as dbm
    from soc_platform.core.models import ActionRequest, Case, CaseEntity, Entity, Evidence
    from soc_platform.domains.phishing.models import Submission
    from soc_platform.domains.vulnerability.models import ConsolidatedFinding, RemediationCampaign
    from soc_platform.intelligence.models import Insight

    with dbm.get_database().session() as s:
        cases = {c.id: c for c in s.execute(select(Case)).scalars()}
        entities = set(s.execute(select(Entity.id)).scalars())
        campaigns = set(s.execute(select(RemediationCampaign.id)).scalars())
        bad = []
        for ln in s.execute(select(CaseEntity)).scalars():
            if ln.case_id not in cases or ln.entity_id not in entities:
                bad.append(f"case link {ln.case_id}->{ln.entity_id}")
        for ev in s.execute(select(Evidence)).scalars():
            if ev.case_id not in cases:
                bad.append(f"evidence {ev.id} of missing case")
        for a in s.execute(select(ActionRequest)).scalars():
            if a.case_id and a.case_id not in cases and a.case_id not in campaigns:
                bad.append(f"action {a.id} of missing case/campaign {a.case_id}")
            for linked in (a.result or {}).get("linked_cases", []):
                if linked not in cases:
                    bad.append(f"action {a.id} linked to missing case {linked}")
        for i in s.execute(select(Insight)).scalars():
            bad += [f"insight {i.id} -> missing entity {e}" for e in i.entity_ids or [] if e not in entities]
        for sub in s.execute(select(Submission)).scalars():
            if sub.case_id and sub.case_id not in cases:
                bad.append(f"submission {sub.id} -> missing case")
        for f in s.execute(select(ConsolidatedFinding)).scalars():
            if f.campaign_id and f.campaign_id not in campaigns:
                bad.append(f"finding {f.id} -> missing campaign")
        for c in cases.values():                            # every citation in an assessment points at its evidence
            a = c.assessment or {}
            idx = set(a.get("evidence_index") or {})
            for cl in a.get("claims") or []:
                if idx and not set(cl["evidence_ids"]) <= idx:
                    bad.append(f"case {c.id} cites unknown evidence {cl['evidence_ids']}")
            if c.created_at.tzinfo is None:
                bad.append(f"case {c.id} naive timestamp")
        assert not bad, bad[:20]


# --------------------------------------------------------------------------------------------- 6. self-check (in product)
@pytest.mark.parametrize("estate", ESTATES)
def test_self_check_passes_on_a_consistent_platform_and_catches_corruption(estate, estate_configs):
    from soc_platform.connectors.registry import ConnectorRegistry
    from soc_platform.core.db import Database
    from soc_platform.core.models import Case, Evidence
    from soc_platform.core.selfcheck import raise_or_resolve, run_self_check
    from soc_platform.intelligence.models import Insight

    db = Database("sqlite://")
    db.create_all()
    with estate_env(estate_configs[estate]) as cfg, db.session() as s:
        _estate(s, ConnectorRegistry.all_fake(), cfg)
        r = run_self_check(s)
        assert r["ok"] and r["passed"] == r["total"] >= 14, [c for c in r["checks"] if not c["ok"]]

        # corrupt it the ways real deployments drift: a second case for one report, orphaned evidence
        ph = s.query(Case).filter(Case.domain == "phishing").first()
        dup = Case(domain="phishing", title=ph.title, severity=ph.severity, status="investigating", attributes=dict(ph.attributes))
        s.add(dup)
        s.add(Evidence(case_id="gone-case", dimension="email", source_tool="x", summary="orphan"))
        s.flush()
        bad = run_self_check(s)
        failing = {c["check"] for c in bad["checks"] if not c["ok"]}
        assert {"one case per reported email", "evidence belongs to a case"} <= failing, failing
        raise_or_resolve(s, bad)
        alert = s.query(Insight).filter(Insight.rule == "platform_integrity").one()
        assert alert.status == "new" and "checks failing" in alert.title

        # repaired -> the finding resolves itself
        s.delete(dup)
        s.query(Evidence).filter(Evidence.case_id == "gone-case").delete()
        s.flush()
        ok = run_self_check(s)
        raise_or_resolve(s, ok)
        assert ok["ok"] and s.query(Insight).filter(Insight.rule == "platform_integrity").one().status == "resolved"


def test_self_check_endpoint_is_for_all_domain_auditors(api):
    c, H = api["c"], api["H"]
    org = api["cfg"]["org"]
    r = J(c.get("/api/v1/admin/self-check", headers=H(f"audrey@{org}", "auditor")))
    assert r["ok"], [x for x in r["checks"] if not x["ok"]]
    assert c.get("/api/v1/admin/self-check", headers=H(f"pia@{org}", "analyst", "phishing")).status_code == 403
    assert c.get("/api/v1/admin/self-check").status_code == 401


def test_budget_alert_and_confirmed_self_check_alerts():
    from soc_platform.config import Settings
    from soc_platform.core.db import Database
    from soc_platform.core.models import LLMCall
    from soc_platform.core.selfcheck import confirm, llm_budget_alert, run_self_check
    from soc_platform.intelligence.models import Insight

    db = Database("sqlite://")
    db.create_all()
    with db.session() as s:
        st = Settings(llm_provider="openai_compatible", llm_monthly_token_budget=1000)
        assert llm_budget_alert(s, st)["alert"] is False and s.query(Insight).count() == 0
        s.add(LLMCall(workflow="t", provider="x", model="m", prompt_redacted="", response="", prompt_tokens=700,
                      completion_tokens=150, status="ok", grounded=True))
        s.flush()
        llm_budget_alert(s, st)
        a = s.query(Insight).filter(Insight.rule == "llm_budget").one()
        assert a.severity == "medium" and "85%" in a.title
        s.add(LLMCall(workflow="t", provider="x", model="m", prompt_redacted="", response="", prompt_tokens=200,
                      completion_tokens=0, status="ok", grounded=True))
        s.flush()
        llm_budget_alert(s, st)
        assert a.severity == "high" and "exhausted" in a.title
        assert llm_budget_alert(s, Settings()) is None                                # no LLM configured: nothing to watch

        # a check that fails once but not on the confirming re-run never alerts
        first = run_self_check(s)
        flaky = {**first, "ok": False, "checks": [{**first["checks"][0], "ok": False}] + first["checks"][1:]}
        assert confirm(s, flaky, wait=0)["ok"] is True
