"""The organisation's report style: brand, logo, Word / PowerPoint templates and a writing guide, uploaded on the
Reports screen and applied to every generated report - never to a figure."""

from __future__ import annotations

import base64
import io
import json
import zipfile

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from soc_platform.core.auth import Principal, Role
from soc_platform.reporting import design as D
from soc_platform.reporting.style import ReportStyleStore, StyleRejected, preview, use_style

LEAD = Principal("lena@acme-demo.com", "Lena", frozenset({Role.LEAD}))
ADMIN = Principal("ada@acme-demo.com", "Ada", frozenset({Role.ADMIN}))
ANALYST = Principal("al@acme-demo.com", "Al", frozenset({Role.ANALYST}))


def word_template(*, header: str = "NORTHWIND HEADER", template: bool = False, macro: bool = False,
                  headings: bool = True) -> bytes:
    d = Document()
    d.sections[0].header.paragraphs[0].text = header
    d.add_paragraph("Sample body text that must not appear in reports")
    if not headings:                                  # a template without the heading styles reports use
        for name in ("Heading 1", "Heading 2"):
            st = d.styles[name]
            st.element.getparent().remove(st.element)
    buf = io.BytesIO()
    d.save(buf)
    return _repack(buf.getvalue(), template=("wordprocessingml.document.main", "wordprocessingml.template.main")
                   if template else None, macro=macro)


def deck_template(*, wide: bool = True, template: bool = False) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = (Inches(13.333), Inches(7.5)) if wide else (Inches(10), Inches(7.5))
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Template sample slide"
    buf = io.BytesIO()
    prs.save(buf)
    return _repack(buf.getvalue(), template=("presentationml.presentation.main", "presentationml.template.main")
                   if template else None)


def _repack(data: bytes, *, template: tuple[str, str] | None = None, macro: bool = False, dtd: bool = False) -> bytes:
    src, out = zipfile.ZipFile(io.BytesIO(data)), io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for info in src.infolist():
            body = src.read(info)
            if template and info.filename == "[Content_Types].xml":
                body = body.replace(template[0].encode(), template[1].encode())
            if dtd and info.filename == "[Content_Types].xml":
                body = body.replace(b"?>", b'?><!DOCTYPE x [<!ENTITY e "boom">]>', 1)
            z.writestr(info, body)
        if macro:
            z.writestr("word/vbaProject.bin", b"\0" * 16)
    return out.getvalue()


def png(colour=(200, 30, 30), size=(120, 40)) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", size, colour).save(buf, "PNG")
    return buf.getvalue()


def docx_text(data: bytes) -> str:
    d = Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells]
    for s in d.sections:
        parts += [p.text for p in s.header.paragraphs] + [p.text for p in s.footer.paragraphs]
    return "\n".join(parts)


# ---------------------------------------------------------------------------------------------- brand


def test_the_house_style_applies_until_someone_sets_another(session):
    assert ReportStyleStore(session).view()["custom"] is False
    text = docx_text(preview(session, "docx"))
    assert "Agentic SOC" in text and D.CLASSIFICATION in text
    assert D.primary() == D.NAVY and D.STYLE.get() is None          # nothing leaks out of a rendering


def test_brand_settings_reach_every_report(session):
    st = ReportStyleStore(session)
    st.set_brand({"organisation": "Northwind Traders", "primary": "#0b3d2e", "accent": "c2410c", "font": "Arial",
                  "classification": "Restricted - Northwind", "footer": "Prepared by the Northwind SOC"}, LEAD, "rebrand")
    body = preview(session, "docx")
    d = Document(io.BytesIO(body))
    text = docx_text(body)
    assert "Northwind Traders" in text and "Restricted - Northwind" in text and "Prepared by the Northwind SOC" in text
    assert d.styles["Normal"].font.name == "Arial"
    assert "0B3D2E" in d.element.xml or any("0B3D2E" in t._tbl.xml for t in d.tables)   # the cover band's colour
    deck = Presentation(io.BytesIO(preview(session, "pptx")))
    slide_text = " ".join(sh.text_frame.text for s in deck.slides for sh in s.shapes if sh.has_text_frame)
    assert "Northwind Traders" in slide_text and "Restricted - Northwind" in slide_text
    assert D.STYLE.get() is None


def test_every_bad_brand_setting_is_named(session):
    with pytest.raises(StyleRejected) as e:
        ReportStyleStore(session).set_brand({"primary": "navy", "font": "Arial; drop", "classification": "",
                                             "organisation": "x" * 81, "colour": "red", "guide": "g" * 9000}, LEAD)
    for needle in ("primary", "font", "classification: cannot be empty", "organisation", "unknown setting 'colour'",
                   "guide: at most"):
        assert needle in str(e.value), needle


def test_only_a_lead_or_an_administrator_changes_the_style_and_every_change_is_audited(session):
    from soc_platform.core.models import AuditRecord

    st = ReportStyleStore(session)
    for who in (ANALYST, Principal("svc", "svc", frozenset({Role.ADMIN}), is_service=True)):
        with pytest.raises(PermissionError):
            st.set_brand({"organisation": "X"}, who)
        with pytest.raises(PermissionError):
            st.upload("logo.png", png(), who)
    st.set_brand({"organisation": "Northwind"}, LEAD, "first")
    st.set_brand({"organisation": "Northwind Traders"}, ADMIN, "second")
    with pytest.raises(StyleRejected, match="changes nothing"):
        st.set_brand({"organisation": "Northwind Traders"}, ADMIN)
    events = [a for a in session.query(AuditRecord).all() if a.event_type == "report_style.changed"]
    assert [e.actor_id for e in events] == [LEAD.id, ADMIN.id]
    assert events[-1].payload["changes"]["organisation"] == {"from": "Northwind", "to": "Northwind Traders"}
    assert [h["note"] for h in st.view()["history"]] == ["second", "first"]


# ---------------------------------------------------------------------------------------------- files


def test_a_word_template_is_the_base_of_every_word_report(session):
    st = ReportStyleStore(session)
    r = st.upload("northwind.dotx", word_template(template=True, headings=False), LEAD)   # a .dotx, no heading styles
    assert r["applied"][0]["slot"] == "word"
    text = docx_text(preview(session, "docx"))
    assert "NORTHWIND HEADER" in text                                  # the template's own header is kept
    assert "Sample body text" not in text                              # its sample content is not
    assert "Report style preview" in text and "A section" in text      # the report itself is there, headings and all


def test_a_powerpoint_template_supplies_the_master_and_must_be_widescreen(session):
    st = ReportStyleStore(session)
    with pytest.raises(StyleRejected, match="16:9"):
        st.upload("old.pptx", deck_template(wide=False), LEAD)
    st.upload("northwind.potx", deck_template(template=True), LEAD)
    deck = Presentation(io.BytesIO(preview(session, "pptx")))
    titles = [sh.text_frame.text for s in deck.slides for sh in s.shapes if sh.has_text_frame]
    assert len(deck.slides) == 2 and "Template sample slide" not in titles   # its sample slide is gone


def test_a_logo_is_re_encoded_and_placed_on_the_cover_header_and_title_slide(session):
    st = ReportStyleStore(session)
    jpeg = io.BytesIO()
    from PIL import Image

    Image.new("RGB", (300, 100), (0, 80, 160)).save(jpeg, "JPEG")
    st.upload("logo.jpg", jpeg.getvalue(), LEAD)
    assert st.active()["files"]["logo"]["detail"] == "300 x 100 px"
    with use_style(session) as s:
        assert s["logo_bytes"].startswith(b"\x89PNG")                    # stored as a clean PNG
    d = Document(io.BytesIO(preview(session, "docx")))
    assert len(d.inline_shapes) >= 1 and d.sections[0].header._element.xpath(".//w:drawing")
    deck = Presentation(io.BytesIO(preview(session, "pptx")))
    assert any(sh.shape_type == 13 for sh in deck.slides[0].shapes)       # a picture on the title slide
    with pytest.raises(StyleRejected, match="PNG or JPEG"):
        st.upload("logo.png", b"\x89PNG not really", LEAD)


def test_a_tiny_png_claiming_gigapixels_is_refused_not_decoded(session):
    import struct
    import zlib

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

    bomb = (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 30_000, 30_000, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(b"")) + chunk(b"IEND", b""))
    assert len(bomb) < 100
    with pytest.raises(StyleRejected, match="at most 5000 pixels"):
        ReportStyleStore(session).upload("logo.png", bomb, LEAD)
    assert ReportStyleStore(session).active_row() is None


@pytest.mark.parametrize("name,data,needle", [
    ("macro.docx", None, "macros"),
    ("macro.docm", b"x", "macro-enabled"),
    ("dtd.docx", None, "document type"),
    ("not-office.docx", b"PK\x03\x04 broken", "not a readable"),
    ("notes.exe", b"MZ", "not a style file"),
    ("huge.zip", None, "too large"),
])
def test_unusable_or_hostile_files_are_refused_with_a_reason(session, name, data, needle):
    if name == "macro.docx":
        data = word_template(macro=True)
    elif name == "dtd.docx":
        data = _repack(word_template(), dtd=True)
    elif name == "huge.zip":
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("brand.json", b" " * 61_000_000)                   # compresses to almost nothing
        data = buf.getvalue()
    with pytest.raises(StyleRejected, match=needle):
        ReportStyleStore(session).upload(name, data, LEAD)
    assert ReportStyleStore(session).active_row() is None                 # nothing half-applied


def test_a_style_pack_applies_everything_it_holds_in_one_version(session):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("northwind/SKILL.md", "# Northwind report voice\nWrite for executives. Say 'incident', not 'breach'.")
        z.writestr("northwind/brand.json", json.dumps({"organisation": "Northwind", "primary": "0B3D2E"}))
        z.writestr("northwind/logo.png", png())
        z.writestr("northwind/report.docx", word_template())
        z.writestr("__MACOSX/._report.docx", b"junk")
    st = ReportStyleStore(session)
    r = st.upload("northwind-style.zip", buf.getvalue(), LEAD, "style pack from brand team")
    assert [a["slot"] for a in r["applied"]] == ["brand", "guide", "logo", "word"]
    doc = st.active()
    assert doc["organisation"] == "Northwind" and "Say 'incident'" in doc["guide"]
    assert set(doc["files"]) == {"logo", "word"} and len(st.view()["history"]) == 1
    st.remove("word", LEAD)
    assert set(st.active()["files"]) == {"logo"}
    st.reset(LEAD)
    assert st.view()["custom"] is False


def test_style_files_live_in_the_database_so_every_replica_renders_with_them(session):
    from soc_platform.core.blobstore import Blob

    ReportStyleStore(session).upload("logo.png", png(), LEAD)
    ref = ReportStyleStore(session).active()["files"]["logo"]["ref"]
    assert ref.startswith("db:report_style/") and session.get(Blob, ref[3:]) is not None


# ---------------------------------------------------------------------------------------------- the narrative


class Writer:
    name = "fake"

    def __init__(self) -> None:
        self.prompts: list[str] = []

    def complete(self, system, user, *, tier, max_tokens=None, model=None):
        from soc_platform.llm.gateway import Completion

        self.prompts.append(user)
        return Completion(json.dumps({"summary": "ok", "claims": [
            {"text": "Exposure is stable.", "kind": "fact", "evidence_ids": ["F1"]}]}), 10, 10, "fake")


def test_the_writing_guide_shapes_model_text_only(session, tmp_path):
    from soc_platform.config import Settings
    from soc_platform.llm.gateway import LLMGateway
    from soc_platform.reporting.builder import build_report

    ReportStyleStore(session).set_brand({"guide": "Write for executives. British English."}, LEAD)
    w = Writer()
    gw = LLMGateway(session, Settings(llm_provider="fake", llm_monthly_token_budget=10**9), provider=w)
    spec = {"title": "Board brief", "audience": "board", "days": 30, "format": "docx",
            "sections": [{"source": "overview", "title": "Overview", "instruction": "summarise"}]}
    with_model = build_report(session, None, spec, tmp_path, llm=gw, by=LEAD.id)
    assert "HOUSE WRITING STYLE" in w.prompts[0] and "British English" in w.prompts[0]
    assert "never changes a fact, a figure" in w.prompts[0]
    without = build_report(session, None, spec, tmp_path, llm=None, by=LEAD.id)       # deterministic text unchanged
    assert with_model["sections"][0]["facts"] == without["sections"][0]["facts"]       # figures identical either way


# ---------------------------------------------------------------------------------------------- the API


def test_the_reports_screen_uploads_previews_and_resets(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    from soc_platform.api import app as appmod
    from soc_platform.config import get_settings
    from soc_platform.core import db as dbm
    from soc_platform.core.auth import issue_dev_token

    secret = "style-secret-0123456789abcdef0123456789"
    for k, v in {"SOC_AUTH_MODE": "dev", "SOC_DEV_JWT_SECRET": secret, "SOC_EMBEDDED_SCHEDULER": "0",
                 "SOC_DATABASE_URL": f"sqlite:///{(tmp_path / 's.db').as_posix()}",
                 "SOC_REPORT_OUTPUT_DIR": str(tmp_path / "rep")}.items():
        monkeypatch.setenv(k, v)
    get_settings.cache_clear()
    dbm._default = None
    appmod.registry.cache_clear()
    try:
        c = TestClient(appmod.app)
        h = lambda u, r: {"Authorization": "Bearer " + issue_dev_token(secret, u, [r])}
        lead, analyst, auditor = h(LEAD.id, "lead"), h(ANALYST.id, "analyst"), h("au@acme-demo.com", "auditor")
        up = {"content_b64": base64.b64encode(png()).decode(), "filename": "logo.png", "note": "brand refresh"}
        assert c.post("/api/v1/reports/style/upload", headers=analyst, json=up).status_code == 403
        assert c.post("/api/v1/reports/style/upload", headers=lead, json=up).status_code == 200
        bad = c.post("/api/v1/reports/style/upload", headers=lead, json={**up, "filename": "logo.gif"})
        assert bad.status_code == 400 and "not a style file" in bad.json()["detail"]
        assert c.post("/api/v1/reports/style", headers=lead, json={"organisation": "Northwind"}).status_code == 200
        view = c.get("/api/v1/reports/style", headers=auditor).json()
        assert view["style"]["organisation"] == "Northwind" and "ref" not in view["style"]["files"]["logo"]
        r = c.get("/api/v1/reports/style/preview?format=docx", headers=auditor)
        assert r.status_code == 200 and "Northwind" in docx_text(r.content)
        assert c.get("/api/v1/reports/style/preview?format=pdf", headers=auditor).status_code == 422
        assert c.post("/api/v1/reports/style/reset", headers=lead, json={}).status_code == 200
        assert c.get("/api/v1/reports/style", headers=auditor).json()["custom"] is False
    finally:
        get_settings.cache_clear()
        dbm._default = None
        appmod.registry.cache_clear()
