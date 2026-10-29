"""Watch lists: block and exception lists of users, hosts, IPs, domains, URLs and hashes, uploaded as Excel or CSV,
shown on every case that involves a listed item - never changing a figure, verdict or action."""

from __future__ import annotations

import base64
import io
import os
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import pytest

from soc_platform.core.watchlists import normalise, parse, read_rows

ROOT = Path(__file__).resolve().parents[2]
SAMPLE = Path(__file__).with_name("data") / "watchlist_sample.xlsx"   # written by openpyxl, as Excel writes them


def xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    """A workbook the way Excel writes one: shared strings, cell references, one part per sheet."""
    strings: list[str] = []

    def sst(v: str) -> int:
        if v not in strings:
            strings.append(v)
        return strings.index(v)

    ns = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    rel = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    parts, sheet_tags, rels = {}, [], []
    for n, (title, rows) in enumerate(sheets.items(), 1):
        body = []
        for ri, row in enumerate(rows, 1):
            cells = []
            for ci, v in enumerate(row):
                ref = f"{chr(65 + ci)}{ri}"
                if isinstance(v, (int, float)):
                    cells.append(f'<c r="{ref}"><v>{v}</v></c>')
                elif v:
                    cells.append(f'<c r="{ref}" t="s"><v>{sst(str(v))}</v></c>')
            body.append(f'<row r="{ri}">{"".join(cells)}</row>')
        parts[f"xl/worksheets/sheet{n}.xml"] = f'<worksheet {ns}><sheetData>{"".join(body)}</sheetData></worksheet>'
        sheet_tags.append(f'<sheet name="{escape(title)}" sheetId="{n}" r:id="rId{n}"/>')
        rels.append(f'<Relationship Id="rId{n}" Type="x/worksheet" Target="worksheets/sheet{n}.xml"/>')
    parts["xl/workbook.xml"] = f'<workbook {ns} {rel}><sheets>{"".join(sheet_tags)}</sheets></workbook>'
    parts["xl/_rels/workbook.xml.rels"] = ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
                                           f'relationships">{"".join(rels)}</Relationships>')
    parts["xl/sharedStrings.xml"] = (f'<sst {ns}>' + "".join(f'<si><t xml:space="preserve">{escape(s)}</t></si>'
                                                             for s in strings) + "</sst>")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for k, v in parts.items():
            z.writestr(k, v)
    return buf.getvalue()


# ---------------------------------------------------------------------------------------------- reading files


def test_identifiers_are_normalised_to_one_form_per_kind():
    assert normalise("ACME\\Jane.Doe") == ("jane.doe", "user")
    assert normalise("Jane Doe <Jane.Doe@Example.org>") == ("jane.doe@example.org", "email")
    assert normalise("10.0.0.1") == ("10.0.0.1", "ip")
    assert normalise("2001:DB8::1") == ("2001:db8::1", "ip")
    assert normalise("HTTPS://Bad.Example/Path/") == ("https://bad.example/path", "url")
    assert normalise("A" * 64) == ("a" * 64, "hash")
    assert normalise("WEB01") == ("web01", "any")                 # a user name or a host: matches either
    assert normalise("web01.corp.example.", "host") == ("web01.corp.example", "host")
    for bad, kind in [("SYSTEM", "user"), ("HOST01$", "user"), ("not an ip", "ip"), ("x y", None), ("a@", "email"),
                      ("", None), ("bad\x01value", None)]:
        assert isinstance(normalise(bad, kind), str), bad


def test_an_excel_workbook_written_by_a_spreadsheet_library_is_read_sheet_by_sheet():
    entries, problems = parse(read_rows(SAMPLE.read_bytes(), SAMPLE.name))
    got = {e["value"]: (e["kind"], e["note"], e["sheet"]) for e in entries}
    assert got["svc.backup"] == ("user", "service account - noisy", "People")
    assert got["pat.example@example.org"][0] == "email"
    assert got["203.0.113.7"][0] == "ip" and got["bad.example"][0] == "domain"
    assert got["https://bad.example/login"][0] == "url" and got["ab" * 32][0] == "hash"
    assert got["scan-01"] == ("host", "vuln scanner", "Indicators")
    assert got["12345"][0] == "any"                               # a number cell, read without ".0"
    assert {p["value"] for p in problems} == {"SYSTEM", "x y"}
    assert all(p["where"] and p["problem"] for p in problems)     # each rejected value says where and why


def test_every_identifier_column_is_read_and_a_type_column_types_untyped_values():
    data = xlsx({"Hosts": [["Hostname", "IP address", "Owner", "Reason"], ["build-07", "198.51.100.9", "jane", "CI"]],
                 "Mixed": [["Value", "Type"], ["corp-share", "Host"], ["evil.example", "domain"], ["nope", "Planet"]]})
    entries, problems = parse(read_rows(data, "lists.xlsx"))
    assert {(e["value"], e["kind"]) for e in entries} == {("build-07", "host"), ("198.51.100.9", "ip"),
                                                          ("corp-share", "host"), ("evil.example", "domain")}
    assert next(e for e in entries if e["value"] == "build-07")["note"] == "CI"
    assert [p["problem"] for p in problems] == ["unknown type (user, email, host, ip, domain, url or hash)"]


def test_a_csv_or_plain_list_works_without_a_header_and_repeats_are_reported():
    data = "# exceptions\njdoe;approved tester\n\nsvc-scan\njdoe;again\n".encode("utf-8-sig")
    entries, problems = parse(read_rows(data, "list.csv"))
    assert [(e["value"], e["note"]) for e in entries] == [("jdoe", "approved tester"), ("svc-scan", "")]
    assert problems[0]["where"] == "line 5" and "repeated" in problems[0]["problem"]


def test_unreadable_or_hostile_files_are_refused_with_a_reason():
    with pytest.raises(ValueError, match="97-2003"):
        read_rows(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64, "old.xls")
    with pytest.raises(ValueError, match="larger than"):
        read_rows(b"a\n" * 3_000_000, "big.csv")
    hostile = io.BytesIO()
    with zipfile.ZipFile(hostile, "w") as z:          # an entity declaration (XXE / "billion laughs")
        z.writestr("xl/workbook.xml", '<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><workbook>&e;</workbook>')
    with pytest.raises(ValueError, match="document type"):
        read_rows(hostile.getvalue(), "x.xlsx")
    with pytest.raises(ValueError, match="readable"):
        read_rows(b"PK\x03\x04 not really a zip", "x.xlsx")


# ---------------------------------------------------------------------------------------------- on cases (API)

SECRET = "test-secret-0123456789abcdef0123456789"
_ENV = ("SOC_AUTH_MODE", "SOC_DEV_JWT_SECRET", "SOC_DATABASE_URL", "SOC_CONNECTORS_CONFIG", "SOC_REPORT_OUTPUT_DIR",
        "SOC_RAW_PAYLOAD_DIR", "SOC_ORG_DOMAINS", "SOC_ENVIRONMENT")


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    from soc_platform.config import get_settings
    from soc_platform.core import db as dbm

    saved = {k: os.environ.get(k) for k in _ENV}

    def restore():
        for k, v in saved.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v
        get_settings.cache_clear()
        dbm._default = None

    tmp = tmp_path_factory.mktemp("wl")
    os.environ.update({"SOC_AUTH_MODE": "dev", "SOC_DEV_JWT_SECRET": SECRET, "SOC_ENVIRONMENT": "test",
                       "SOC_DATABASE_URL": f"sqlite:///{tmp / 'wl.db'}", "SOC_ORG_DOMAINS": "acme-demo.com",
                       "SOC_CONNECTORS_CONFIG": str(ROOT / "config" / "connectors.yaml"),
                       "SOC_REPORT_OUTPUT_DIR": str(tmp / "reports"), "SOC_RAW_PAYLOAD_DIR": str(tmp / "raw")})
    get_settings.cache_clear()
    dbm._default = None
    from fastapi.testclient import TestClient

    from soc_platform.api import app as appmod

    appmod.registry.cache_clear()
    try:
        c = TestClient(appmod.app)
        a = tok(c, "seed", "analyst")
        assert c.post("/api/v1/incidents/run", headers=a).status_code == 200
        assert c.post("/api/v1/phishing/ingest", headers=a).status_code == 200
        yield c
    finally:
        appmod.registry.cache_clear()
        restore()


def tok(client, user, roles, domains=""):
    t = client.get(f"/api/v1/dev/token?user={user}&roles={roles}&domains={domains}").json()["token"]
    return {"Authorization": f"Bearer {t}"}


def upload(client, who, list_name, data, **kw):
    return client.post(f"/api/v1/watchlists/{list_name}/upload", headers=who,
                       json={"content_b64": base64.b64encode(data).decode(), "filename": kw.pop("filename", "list.xlsx"), **kw})


def _world(client):
    """Identifiers of the estate's own cases, read from the cases (no sample names hard-coded)."""
    h = tok(client, "reader", "analyst")
    views = [client.get(f"/api/v1/cases/{c['id']}", headers=h).json() for c in client.get("/api/v1/cases", headers=h).json()]
    person = next(e for v in views for e in v["entities"] if e["kind"] == "identity" and v["case"]["domain"] == "incident")
    host = next(e for v in views for e in v["entities"] if e["kind"] == "asset" and e["attributes"].get("hostname"))
    domain = next(e["name"].split(":", 1)[1] for v in views for e in v["entities"]
                  if e["kind"] == "indicator" and e["name"].startswith("domain:") and e["name"].count(".") >= 2)
    sender = next(e["name"].split(":", 1)[1] for v in views for e in v["entities"]
                  if e["kind"] == "indicator" and e["role"] == "sender"
                  and e["name"].split("@")[-1] not in ("acme-demo.com",))
    return views, person, host, domain, sender


def _matches(client, h, case_id):
    return client.get(f"/api/v1/cases/{case_id}", headers=h).json()["watchlists"]


def test_a_case_says_which_listed_users_hosts_and_indicators_it_involves(client):
    views, person, host, domain, _ = _world(client)
    lead, analyst = tok(client, "lena", "lead"), tok(client, "ana", "analyst")
    parent = domain.split(".", 1)[1]                    # listing a domain covers its sub-domains
    before = {v["case"]["id"]: (v["case"]["severity"], v["case"]["verdict"], v["case"]["confidence"],
                                sorted((a["action_type"], a["status"]) for a in v["actions"])) for v in views}
    local = person["name"].split("@")[0]
    book = xlsx({"Users": [["Username", "Reason"], [f"CORP\\{local}", "under review"]],
                 "Hosts": [["Hostname", "Notes"], [host["attributes"]["hostname"].upper(), "legacy"]],
                 "IOCs": [["Value", "Type"], [parent, "Domain"]]})
    preview = upload(client, lead, "block", book, dry_run=True)
    assert preview.status_code == 200 and len(preview.json()["added"]) == 3 and preview.json()["dry_run"]
    assert client.get("/api/v1/watchlists", headers=analyst).json()["lists"][0]["count"] == 0   # preview changes nothing

    r = upload(client, lead, "block", book, note="quarterly review")
    assert r.status_code == 200 and r.json()["count_after"] == 3
    found = {}
    for v in views:
        for m in _matches(client, analyst, v["case"]["id"]):
            found.setdefault(m["entry"]["kind"], set()).add(m["entity_id"])
            assert m["list"] == "block" and m["entry"]["added_by"] == "lena" and m["matched_on"]
    assert person["id"] in found["user"] and host["id"] in found["host"] and found["domain"]
    # the case list carries the badge; nothing decisional moved
    listed = {c["id"]: c for c in client.get("/api/v1/cases", headers=analyst).json()}
    for v in views:
        cid = v["case"]["id"]
        assert ("block" in listed[cid]["watchlists"]) == bool(_matches(client, analyst, cid))
        now = client.get(f"/api/v1/cases/{cid}", headers=analyst).json()
        assert (now["case"]["severity"], now["case"]["verdict"], now["case"]["confidence"],
                sorted((a["action_type"], a["status"]) for a in now["actions"])) == before[cid]


def test_a_bare_user_name_never_matches_an_outside_address_that_shares_it(client):
    *_, sender = _world(client)
    lead, analyst = tok(client, "lena", "lead"), tok(client, "ana", "analyst")
    local = sender.split("@")[0]
    assert upload(client, lead, "exception", f"{local}\n".encode(), filename="x.csv").status_code == 200
    hits = [m for c in client.get("/api/v1/cases", headers=analyst).json()
            for m in _matches(client, analyst, c["id"]) if m["entry"]["value"] == local]
    assert all(m["kind"] != "indicator" for m in hits)          # the outside sender is not the listed person
    assert upload(client, lead, "exception", sender.encode(), filename="x.csv").status_code == 200
    hits = [m for c in client.get("/api/v1/cases", headers=analyst).json()
            for m in _matches(client, analyst, c["id"]) if m["entry"]["value"] == sender]
    assert hits and all(m["list"] == "exception" and m["matched_on"].startswith("sender address") for m in hits)


def test_only_a_lead_changes_the_lists_and_every_change_is_audited(client):
    lead, analyst, admin = tok(client, "lena", "lead"), tok(client, "ana", "analyst"), tok(client, "adam", "admin")
    data = b"svc-test-01\n"
    assert upload(client, analyst, "block", data).status_code == 403
    assert upload(client, admin, "block", data).status_code == 403
    assert upload(client, lead, "nonsense", data).status_code == 400
    assert upload(client, lead, "block", b"Username\nSYSTEM\n").status_code == 400   # nothing listable
    r = upload(client, lead, "block", data, filename="svc.csv", note="noisy service")
    assert r.status_code == 200
    # the same value on the other list is refused, not silently duplicated
    clash = upload(client, lead, "exception", data).json()
    assert "already on the block list" in clash["detail"]
    entry = next(e for e in client.get("/api/v1/watchlists", headers=analyst).json()["lists"][0]["entries"]
                 if e["value"] == "svc-test-01")
    assert entry["source"] == "svc.csv" and entry["note"] == "noisy service"
    assert client.delete(f"/api/v1/watchlists/entries/{entry['id']}?reason=retired", headers=analyst).status_code == 403
    assert client.delete(f"/api/v1/watchlists/entries/{entry['id']}?reason=retired", headers=lead).status_code == 200
    assert client.delete(f"/api/v1/watchlists/entries/{entry['id']}", headers=lead).status_code == 404
    audit = client.get("/api/v1/audit?event_type=watchlist.uploaded", headers=tok(client, "aud", "auditor")).json()
    up = next(a for a in audit if "user:svc-test-01" in a["payload"].get("added", []) or
              "any:svc-test-01" in a["payload"].get("added", []))
    assert up["actor_id"] == "lena" and up["payload"]["file"] == "svc.csv"
    gone = client.get("/api/v1/audit?event_type=watchlist.entry_removed", headers=tok(client, "aud", "auditor")).json()
    assert any(a["payload"]["value"] == "svc-test-01" and a["payload"]["reason"] == "retired" for a in gone)


def test_replacing_a_list_removes_what_the_file_no_longer_has_and_cases_follow_at_once(client):
    _, _, host, *_ = _world(client)
    lead, analyst = tok(client, "lena", "lead"), tok(client, "ana", "analyst")
    hostname = host["attributes"]["hostname"].lower()
    assert upload(client, lead, "block", xlsx({"S": [["Hostname"], [hostname]]})).status_code == 200
    cases = [c["id"] for c in client.get("/api/v1/cases", headers=analyst).json()]
    assert any(m["entity_id"] == host["id"] for c in cases for m in _matches(client, analyst, c))
    r = upload(client, lead, "block", xlsx({"S": [["Hostname"], ["unrelated-host-99"]]}), mode="replace").json()
    assert hostname in r["removed"] and r["count_after"] == 1
    assert not any(m["entity_id"] == host["id"] and m["list"] == "block"
                   for c in cases for m in _matches(client, analyst, c))
    csv_text = client.get("/api/v1/watchlists/block/export", headers=analyst).text
    assert "unrelated-host-99" in csv_text and hostname not in csv_text


def test_the_investigation_record_lists_the_watch_list_matches(client):
    from docx import Document

    _, person, *_ = _world(client)
    lead, analyst = tok(client, "lena", "lead"), tok(client, "ana", "analyst")
    assert upload(client, lead, "block", person["name"].encode(), filename="p.csv").status_code == 200
    cid = next(c["id"] for c in client.get("/api/v1/cases", headers=analyst).json() if "block" in c["watchlists"])
    r = client.get(f"/api/v1/cases/{cid}/report", headers=analyst)
    assert r.status_code == 200
    text = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    cells = " ".join(c.text for t in Document(io.BytesIO(r.content)).tables for row in t.rows for c in row.cells)
    assert "Watch lists" in text and "Block list" in cells and person["name"] in cells
