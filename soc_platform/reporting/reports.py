"""Recurring and per-case reporting (VM-F13, VM-F14, VM-T08, IM-F12, PH-F14, U01, U05, U15).

All figures are computed deterministically by the domain services; this module only
lays them out. Narrative commentary comes from the LLM gateway (grounded on the
computed figures, which are passed as evidence) or from deterministic sentences when
no model is configured. the client's own templates can be supplied: a ``.docx`` template is
used as the base document (its styles/header/footer are kept) and a ``.pptx``
template supplies the slide master.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt, RGBColor
from pptx import Presentation
from pptx.util import Inches
from pptx.util import Pt as PPt
from sqlalchemy.orm import Session

from soc_platform.core.models import utcnow
from soc_platform.domains.vulnerability.models import ReportRun
from soc_platform.llm.gateway import LLMGateway


def _commentary(llm: LLMGateway | None, workflow: str, question: str, facts: dict[str, Any]) -> tuple[str, str]:
    evidence = [{"id": f"E{i + 1}", "claim": f"{k}: {v}", "source": "computed"} for i, (k, v) in enumerate(facts.items())]
    if llm is not None:
        g = llm.grounded(workflow, question, evidence, tier="small")
        if g.get("source") == "llm" and g.get("claims"):
            return " ".join(c["text"] for c in g["claims"]), "llm (grounded on computed figures)"
    return _deterministic_commentary(facts), "deterministic"


def _deterministic_commentary(f: dict[str, Any]) -> str:
    parts = []
    if "open" in f:
        parts.append(f"{f['open']} open findings")
    if f.get("kev_open"):
        parts.append(f"{f['kev_open']} on the CISA KEV list")
    if f.get("sla_breached"):
        parts.append(f"{f['sla_breached']} past SLA")
    if f.get("internet_exposed_open"):
        parts.append(f"{f['internet_exposed_open']} on internet-exposed assets")
    s = ("Exposure summary: " + ", ".join(parts) + ".") if parts else ""
    if f.get("by_priority"):
        s += " Priority mix " + ", ".join(f"{k}={v}" for k, v in sorted(f["by_priority"].items())) + "."
    return s or "No material change."


def _doc(template: str | Path | None) -> Document:
    d = Document(str(template)) if template and Path(template).exists() else Document()
    if not (template and Path(template).exists()):
        d.styles["Normal"].font.name = "Calibri"
        d.styles["Normal"].font.size = Pt(10.5)
    return d


def _table(doc: Document, header: list[str], rows: list[list[Any]]) -> None:
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1" if "Light Grid Accent 1" in [s.name for s in doc.styles] else t.style
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = str(h)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = "" if v is None else str(v)


def _footer(doc: Document, source: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(f"Generated {utcnow():%Y-%m-%d %H:%M} UTC by the SOC platform. All figures computed from connected "
                    f"tools; commentary: {source}. Asset match rate and unavailable sources are stated where relevant.")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


class ReportService:
    def __init__(self, session: Session, out_dir: str | Path, *, llm: LLMGateway | None = None,
                 templates: dict[str, str] | None = None) -> None:
        self.s = session
        self.out = Path(out_dir)
        self.out.mkdir(parents=True, exist_ok=True)
        self.llm = llm
        self.templates = templates or {}

    def _record(self, kind: str, path: Path, metrics: dict[str, Any], by: str) -> ReportRun:
        from soc_platform.core.crypto import seal_file

        seal_file(path)  # reports carry case details: encrypted at rest like raw payloads
        run = ReportRun(kind=kind, path=str(path), metrics=metrics, generated_by=by)
        self.s.add(run)
        self.s.flush()
        return run

    # ------------------------------------------------------------------ VM-F13 daily exposure

    def daily_exposure(self, vm: Any, *, by: str = "system:scheduler") -> ReportRun:
        m = vm.metrics()
        cov = vm.coverage()
        text, src = _commentary(self.llm, "report.daily_exposure", "Write a 3-sentence daily exposure commentary.", m)
        d = _doc(self.templates.get("daily_exposure"))
        d.add_heading(f"Daily Exposure Report - {utcnow():%d %b %Y}", 0)
        d.add_paragraph(text)
        d.add_heading("Headline figures", 1)
        _table(d, ["Measure", "Value"], [["Open findings", m["open"]], ["KEV-listed open", m["kev_open"]],
                                         ["Internet-exposed open", m["internet_exposed_open"]],
                                         ["Past SLA", m["sla_breached"]], ["Remediated (total)", m["remediated"]],
                                         ["Risk accepted", m["risk_accepted"]],
                                         ["Asset match rate", f"{(m['asset_match_rate'] or 0) * 100:.1f}%"]])
        d.add_heading("Priority and ageing", 1)
        _table(d, ["Priority", "Open"], [[k, v] for k, v in sorted(m["by_priority"].items())])
        _table(d, ["Age (days)", "Open"], [[k, m["ageing"].get(k, 0)] for k in ("0-7", "8-30", "31-90", "90+")])
        d.add_heading("Top vulnerabilities", 1)
        _table(d, ["CVE", "Assets"], [[x["cve"], x["assets"]] for x in m["top_cves"]])
        d.add_heading("Coverage and data quality", 1)
        _table(d, ["Check", "Result"], [["Assets in scope", cov["assets"]], ["Missing EDR", ", ".join(cov["missing_edr"]) or "none"],
                                        ["Not in CMDB", ", ".join(cov["not_in_cmdb"]) or "none"],
                                        ["Unresolved identity queue", cov["unresolved_queue"]]])
        _footer(d, src)
        path = self.out / f"daily_exposure_{utcnow():%Y%m%d_%H%M%S}.docx"
        d.save(path)
        return self._record("daily_exposure", path, m, by)

    # ------------------------------------------------------------------ VM-F13 weekly VM report

    def weekly_vm(self, vm: Any, *, by: str = "system:scheduler") -> ReportRun:
        from sqlalchemy import select

        from soc_platform.domains.vulnerability.models import ActionPlan, ConsolidatedFinding, RemediationCampaign

        m = vm.metrics()
        text, src = _commentary(self.llm, "report.weekly_vm", "Write a weekly vulnerability management summary.", m)
        d = _doc(self.templates.get("weekly_vm"))
        d.add_heading(f"Weekly Vulnerability Management Report - week of {utcnow():%d %b %Y}", 0)
        d.add_paragraph(text)
        d.add_heading("Per-team performance", 1)
        _table(d, ["Platform team", "Open", "Past SLA", "Remediated", "Risk accepted"],
               [[t, v.get("open", 0), v.get("sla_breached", 0), v.get("remediated", 0), v.get("risk_accepted", 0)]
                for t, v in sorted(m["per_team"].items())])
        d.add_heading("Remediation campaigns", 1)
        rows = []
        for c in self.s.execute(select(RemediationCampaign)).scalars():
            plans = self.s.execute(select(ActionPlan).where(ActionPlan.campaign_id == c.id)).scalars().all()
            rows.append([c.title, c.status, len(plans), sum(1 for p in plans if p.status == "acknowledged"),
                         sum(1 for p in plans if p.followups), c.target_date.date() if c.target_date else ""])
        _table(d, ["Campaign", "Status", "Teams", "Acknowledged", "Escalated", "Target"], rows or [["none", "", "", "", "", ""]])
        d.add_heading("P1 findings", 1)
        p1 = self.s.execute(select(ConsolidatedFinding).where(ConsolidatedFinding.priority_band == "P1")).scalars().all()
        _table(d, ["CVE", "Asset", "Team", "SLA due", "Status", "Seen by"],
               [[f.cve, f.asset_name, f.platform_team, f.sla_due.date() if f.sla_due else "", f.status,
                 ", ".join(sorted(f.sources))] for f in p1])
        d.add_heading("Trend indicators", 1)
        _table(d, ["Measure", "Value"], [["MTTR (days)", m["mttr_days"]], ["Recurrence rate", m["recurrence_rate"]],
                                         ["Reopened (total)", m["reopened_total"]]])
        _footer(d, src)
        path = self.out / f"weekly_vm_{utcnow():%Y%m%d_%H%M%S}.docx"
        d.save(path)
        return self._record("weekly_vm", path, m, by)

    # ------------------------------------------------------------------ weekly management presentation

    def weekly_management_deck(self, vm: Any, incident: Any | None = None, phishing: Any | None = None, *,
                               by: str = "system:scheduler") -> ReportRun:
        m = vm.metrics()
        tpl = self.templates.get("weekly_mgmt_pptx")
        prs = Presentation(tpl) if tpl and Path(tpl).exists() else Presentation()
        title = prs.slides.add_slide(prs.slide_layouts[0])
        title.shapes.title.text = "SOC Weekly Management Update"
        title.placeholders[1].text = f"Week of {utcnow():%d %b %Y}"

        def bullets(heading: str, lines: list[str]) -> None:
            sl = prs.slides.add_slide(prs.slide_layouts[1])
            sl.shapes.title.text = heading
            tf = sl.placeholders[1].text_frame
            tf.text = lines[0] if lines else ""
            for ln in lines[1:]:
                tf.add_paragraph().text = ln
            for p in tf.paragraphs:
                for r in p.runs:
                    r.font.size = PPt(18)

        bullets("Vulnerability exposure", [(f"Open findings: {m['open']} (P1 {m['by_priority'].get('P1', 0)}, "
                                           f"P2 {m['by_priority'].get('P2', 0)})"),
                                           f"KEV-listed open: {m['kev_open']}; internet-exposed: {m['internet_exposed_open']}",
                                           f"Past SLA: {m['sla_breached']}; MTTR: {m['mttr_days'] or 'n/a'} days",
                                           f"Asset match rate: {(m['asset_match_rate'] or 0) * 100:.1f}%"])
        if incident is not None:
            h = incident.handover(hours=24 * 7)
            bullets("Incidents", [f"Open incidents: {h['open_total']}",
                                  "By severity: " + ", ".join(f"{k} {v}" for k, v in h["open_by_severity"].items()),
                                  f"Awaiting approval: {len(h['awaiting_approval'])}",
                                  f"Closed this week: {h['closed_this_shift']}"])
        if phishing is not None:
            p = phishing.metrics()
            bullets("Reported phishing", [(f"Reports: {p['reported']}; auto-closed: {p['auto_closed']} "
                                          f"(QA-sampled {p['sampled_for_qa']})"),
                                          "Verdicts: " + ", ".join(f"{k} {v}" for k, v in p["verdict_mix"].items()),
                                          f"Campaigns: {p['campaigns']}; repeat clickers: {len(p['repeat_clickers'])}",
                                          f"Median time to containment: {p['time_to_containment_minutes']['median'] or 'n/a'} min"])
        sl = prs.slides.add_slide(prs.slide_layouts[5])
        sl.shapes.title.text = "Per-team vulnerability status"
        rows = sorted(m["per_team"].items())
        tbl = sl.shapes.add_table(len(rows) + 1, 3, Inches(0.5), Inches(1.5), Inches(9), Inches(0.4 * (len(rows) + 1))).table
        for i, h in enumerate(["Team", "Open", "Past SLA"]):
            tbl.cell(0, i).text = h
        for r, (team, v) in enumerate(rows, start=1):
            tbl.cell(r, 0).text, tbl.cell(r, 1).text, tbl.cell(r, 2).text = team, str(v.get("open", 0)), str(v.get("sla_breached", 0))
        path = self.out / f"weekly_mgmt_{utcnow():%Y%m%d_%H%M%S}.pptx"
        prs.save(path)
        return self._record("weekly_mgmt", path, m, by)

    # ------------------------------------------------------------------ IM-F12 / U15 investigation record & evidence pack

    def investigation_report(self, view: dict[str, Any], *, by: str) -> ReportRun:
        c = view["case"]
        d = _doc(self.templates.get("investigation"))
        d.add_heading(f"Investigation record - {c['title']}", 0)
        _table(d, ["Field", "Value"], [["Case", c["id"]], ["Domain", c["domain"]], ["Status", c["status"]],
                                       ["Verdict", c["verdict"]], ["Severity", c["severity"]],
                                       ["Confidence", c["confidence"]], ["Opened", c["created_at"]]])
        d.add_heading("Summary", 1)
        d.add_paragraph(c["summary"] or "")
        comp = view.get("completeness") or {}
        if comp.get("unavailable"):
            d.add_paragraph("INCOMPLETE - unavailable sources: " + ", ".join(u["source"] for u in comp["unavailable"]))
        d.add_heading("Findings (facts)", 1)
        for cl in view["assessment"].get("facts", []):
            d.add_paragraph(f"{cl['text']} [{', '.join(cl['evidence_ids'])}]", style="List Bullet")
        if view["assessment"].get("inferences"):
            d.add_heading("Inferences", 1)
            for cl in view["assessment"]["inferences"]:
                d.add_paragraph(f"{cl['text']} [{', '.join(cl['evidence_ids'])}]", style="List Bullet")
        d.add_heading("MITRE ATT&CK", 1)
        _table(d, ["Technique", "Name", "Basis"], [[t["technique"], t.get("name", ""), t.get("basis", "")]
                                                  for t in view["assessment"].get("mitre", [])])
        d.add_heading("Entities", 1)
        _table(d, ["Kind", "Role", "Name", "Seen by"], [[e["kind"], e["role"], e["name"], ", ".join(e["seen_by"])]
                                                         for e in view["entities"]])
        d.add_heading("Timeline", 1)
        _table(d, ["Time", "Source", "Event"], [[t["ts"], t.get("tool"), t["title"]] for t in view["timeline"]])
        d.add_heading("Evidence", 1)
        for dim, items in view["evidence"].items():
            d.add_heading(dim, 2)
            for it in items:
                d.add_paragraph(f"[{it['type']}] {it['source']}: {it['summary']}" + (f" ({it['deep_link']})" if it.get("deep_link") else ""),
                                style="List Bullet")
        d.add_heading("Actions", 1)
        _table(d, ["Action", "Status", "Approver", "Rationale"], [[a["action_type"], a["status"], a.get("approver") or "",
                                                                   a["rationale"]] for a in view["actions"]])
        d.add_heading("Analyst decisions", 1)
        _table(d, ["Analyst", "Verdict", "Reasoning", "At"], [[x["analyst"], x["verdict"], x["reasoning"], x["at"]]
                                                               for x in view["dispositions"]] or [["", "", "", ""]])
        d.add_heading("Audit trail (chain of custody)", 1)
        _table(d, ["Seq", "Time", "Actor", "Event"], [[a["seq"], a["ts"], a["actor"], a["event"]] for a in view["audit"]])
        _footer(d, "case record")
        path = self.out / f"investigation_{c['id']}_{utcnow():%Y%m%d_%H%M%S}.docx"
        d.save(path)
        return self._record("investigation", path, {"case_id": c["id"]}, by)
