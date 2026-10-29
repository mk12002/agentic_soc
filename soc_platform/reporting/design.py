"""The house style every generated report shares, in Word and in PowerPoint.

Reports are the work clients most want automated, so they must be fit to send as they are: a cover band, headline
figures as tiles, tables with a clear header and readable rows, status in colour, a running header and a footer with
page numbers on every page, and charts with whole-number axes. Layout only - every figure is computed elsewhere and
passed in; nothing here calculates or words a finding.

The organisation's report style (``reporting/style.py``, set on the Reports screen) replaces the house colours, font,
organisation name, classification line and footer, adds its logo, and supplies a Word template (styles, page set-up,
header, footer - the house header / footer are added only where the template has none) and a PowerPoint master. It
is read from ``STYLE`` (a context variable set by ``style.use_style``), so concurrent reports never mix styles;
outside that context the house style applies.
"""

from __future__ import annotations

from contextvars import ContextVar
from typing import Any

from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = "1F2A44"
ACCENT = "2563EB"
TEXT = "1F2937"
MUTED = "6B7280"
ROW_ALT = "F3F6FB"
TILE = "EEF2F8"
BORDER = "D5DCE8"
TONES = {"bad": "B91C1C", "warn": "B45309", "good": "15803D", "neutral": NAVY, "info": ACCENT}
SEVERITY_TONE = {"critical": "bad", "high": "warn", "medium": "warn", "low": "info", "informational": "neutral",
                 "P1": "bad", "P2": "warn", "P3": "info", "P4": "neutral"}
CLASSIFICATION = "Confidential - internal"
MISSING = "not yet measured"

STYLE: ContextVar[dict[str, Any] | None] = ContextVar("report_style", default=None)


def _style() -> dict[str, Any]:
    return STYLE.get() or {}


def primary() -> str:
    return _style().get("primary") or NAVY


def accent() -> str:
    return _style().get("accent") or ACCENT


def font_name() -> str:
    return _style().get("font") or "Calibri"


def classification() -> str:
    return _style().get("classification") or CLASSIFICATION


def brand() -> str:
    """The name in every running header and slide footer: the organisation's, else the platform's."""
    return _style().get("organisation") or "Agentic SOC"


def footer_text() -> str:
    return _style().get("footer") or "Every figure is computed from the connected tools' records"


def logo() -> bytes | None:
    return _style().get("logo_bytes")


def writing_guide() -> str:
    return _style().get("guide") or ""


def tone_colours() -> dict[str, str]:
    return {**TONES, "neutral": primary(), "info": accent()}


def rgb(hexcode: str) -> RGBColor:
    return RGBColor.from_string(hexcode)


def fmt(value: Any, unit: str = "", *, missing: str = MISSING) -> str:
    """One way to show a number everywhere: thousands separators, at most one decimal, a unit, and words - never
    'None' or 'n/a' - when there is no value."""
    if value is None or value == "":
        return missing
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, int | float):
        v = float(value)
        s = f"{int(v):,}" if v.is_integer() else f"{v:,.1f}"
        return f"{s} {unit}".strip() if unit and unit != "%" else (s + "%" if unit == "%" else s)
    return str(value)


# ---------------------------------------------------------------------------------------------------- Word


def _shade(cell: Any, hexcode: str) -> None:
    tc = cell._tc.get_or_add_tcPr()
    for old in tc.findall(qn("w:shd")):
        tc.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcode)
    tc.append(shd)


def _cell_margins(table: Any, top: int = 50, bottom: int = 50, left: int = 90, right: int = 90) -> None:
    tblpr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)


def _borders(table: Any, color: str = BORDER, inside: bool = True) -> None:
    tblpr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right") + (("insideH", "insideV") if inside else ()):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        b.append(el)
    tblpr.append(b)


def _no_borders(table: Any) -> None:
    tblpr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "nil")
        b.append(el)
    tblpr.append(b)


def _field(paragraph: Any, instruction: str) -> None:
    """A Word field (PAGE, NUMPAGES): Word fills it when the document is opened or printed."""
    run = paragraph.add_run()
    begin, text, end = OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(text)
    run._r.append(end)
    run.font.size = Pt(8)
    run.font.color.rgb = rgb(MUTED)


def _run(paragraph: Any, text: str, *, size: float = 10, bold: bool = False, color: str = TEXT,
         italic: bool = False) -> Any:
    r = paragraph.add_run(text)
    r.font.size, r.font.bold, r.font.italic = Pt(size), bold, italic
    r.font.color.rgb = rgb(color)
    return r


def new_document(title: str, *, template: Any = None, landscape: bool = False) -> Any:
    """An A4 document in the house style with a running header (report name, classification) and a footer with
    page X of Y. A client template keeps its own styles, header and footer."""
    import io
    from pathlib import Path

    from docx import Document

    if template and Path(template).exists():
        return Document(str(template))
    if _style().get("word_bytes"):
        return _from_word_template(Document(io.BytesIO(_style()["word_bytes"])), title, landscape)
    d = Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = font_name(), Pt(10)
    st.font.color.rgb = rgb(TEXT)
    st.paragraph_format.space_after = Pt(4)
    for name, size, color in (("Heading 1", 13, primary()), ("Heading 2", 11, primary())):
        h = d.styles[name]
        h.font.name, h.font.size, h.font.bold = font_name(), Pt(size), True
        h.font.color.rgb = rgb(color)
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(14), Pt(5)
        h.paragraph_format.keep_with_next = True
    sec = d.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    else:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.6)
    sec.header_distance = sec.footer_distance = Cm(0.8)
    _house_header(sec, title)
    _house_footer(sec)
    return d


def _house_header(sec: Any, title: str) -> None:
    import io

    head = sec.header.paragraphs[0]
    head.text = ""
    if logo():
        head.add_run().add_picture(io.BytesIO(logo()), height=Cm(0.55))
        head.add_run("  ")
    _run(head, brand() + "  ·  " + title, size=8, color=MUTED)
    _run(head, "\t\t" + classification(), size=8, color=MUTED, bold=True)


def _house_footer(sec: Any) -> None:
    foot = sec.footer.paragraphs[0]
    foot.text = ""
    _run(foot, footer_text() + "  ·  page ", size=8, color=MUTED)
    _field(foot, "PAGE")
    _run(foot, " of ", size=8, color=MUTED)
    _field(foot, "NUMPAGES")


def _has_content(part: Any) -> bool:
    return any(p.text.strip() for p in part.paragraphs) or bool(part._element.xpath(".//w:drawing"))


def _from_word_template(d: Any, title: str, landscape: bool) -> Any:
    """The organisation's template as the base: its sample body removed, its styles, page set-up, header and footer
    kept; the house header / footer only where it has none."""
    from docx.enum.style import WD_STYLE_TYPE

    body = d.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    names = {s.name for s in d.styles}
    for name, size in (("Heading 1", 13), ("Heading 2", 11)):   # headings the reports use, if the template lacks them
        if name not in names:
            h = d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            h.base_style = d.styles["Normal"] if "Normal" in names else None
            h.font.size, h.font.bold, h.font.name = Pt(size), True, font_name()
            h.font.color.rgb = rgb(primary())
            h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(14), Pt(5)
            h.paragraph_format.keep_with_next = True
    sec = d.sections[0]
    if landscape and sec.page_width and sec.page_height and sec.page_width < sec.page_height:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    if not _has_content(sec.header):
        _house_header(sec, title)
    if not _has_content(sec.footer):
        _house_footer(sec)
    return d


def cover(d: Any, title: str, subtitle: str, meta: str) -> None:
    """A full-width navy band with the report name, the period and where the figures come from."""
    t = d.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _no_borders(t)
    _cell_margins(t, 160, 160, 220, 220)
    c = t.rows[0].cells[0]
    _shade(c, primary())
    p = c.paragraphs[0]
    if logo():
        import io

        p.add_run().add_picture(io.BytesIO(logo()), height=Cm(1.1))
        p = c.add_paragraph()
    _run(p, title, size=20, bold=True, color="FFFFFF")
    p2 = c.add_paragraph()
    _run(p2, subtitle, size=11, color="DCE3F0")
    p3 = c.add_paragraph()
    _run(p3, meta, size=8, color="AEB9CF")
    d.add_paragraph()


def tiles(d: Any, items: list[tuple[str, Any, str]], *, unit_map: dict[str, str] | None = None) -> None:
    """Headline figures as a row of tiles: (label, value, tone). Tone colours the value: bad / warn / good /
    neutral / info."""
    if not items:
        return
    t = d.add_table(rows=1, cols=len(items))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _no_borders(t)
    _cell_margins(t, 120, 120, 140, 140)
    for cell, (label, value, tone) in zip(t.rows[0].cells, items, strict=True):
        _shade(cell, TILE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _run(p, fmt(value, (unit_map or {}).get(label, ""), missing="–"), size=20, bold=True,
             color=tone_colours().get(tone, primary()))
        lp = cell.add_paragraph()
        _run(lp, label, size=8, color=MUTED)
    d.add_paragraph()


def table(d: Any, header: list[str], rows: list[list[Any]], *, widths_cm: list[float] | None = None,
          numeric: set[int] | None = None, tone_col: int | None = None, tones: dict[str, str] | None = None,
          empty: str = "Nothing to report.", max_rows: int | None = None, more_note: str = "") -> None:
    """A table in the house style: navy header, alternating rows, thin borders, numbers right-aligned, an optional
    column whose values are coloured by status (``tones`` maps a value to a tone)."""
    if not rows:
        p = d.add_paragraph()
        _run(p, empty, size=9, color=MUTED, italic=True)
        return
    shown = rows[:max_rows] if max_rows else rows
    numeric = numeric or set()
    t = d.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = widths_cm is None
    _borders(t)
    _cell_margins(t)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        _shade(c, primary())
        p = c.paragraphs[0]
        if i in numeric:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _run(p, str(h), size=8.5, bold=True, color="FFFFFF")
    tr = t.rows[0]._tr.get_or_add_trPr()                  # the header repeats on every page
    rep = OxmlElement("w:tblHeader")
    rep.set(qn("w:val"), "true")
    tr.append(rep)
    for r_i, row in enumerate(shown):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            c = cells[i]
            if r_i % 2:
                _shade(c, ROW_ALT)
            p = c.paragraphs[0]
            if i in numeric:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            text = fmt(v, missing="–")
            color, bold = TEXT, False
            if tone_col is not None and i == tone_col:
                tone = (tones or SEVERITY_TONE).get(str(v), "neutral")
                color, bold = tone_colours().get(tone, TEXT), True
            _run(p, text, size=9, color=color, bold=bold)
    if widths_cm:
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        t._tbl.tblPr.append(layout)
        for i, w in enumerate(widths_cm):
            t.columns[i].width = Cm(w)
        for row in t.rows:
            for i, w in enumerate(widths_cm):
                row.cells[i].width = Cm(w)
    if max_rows and len(rows) > max_rows:
        p = d.add_paragraph()
        _run(p, more_note or f"{len(rows) - max_rows} more in the console.", size=8, color=MUTED, italic=True)
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def section(d: Any, title: str, intro: str | None = None) -> None:
    d.add_heading(title, level=1)
    if intro:
        p = d.add_paragraph()
        _run(p, intro, size=9.5, color=MUTED)


def paragraph(d: Any, text: str, *, size: float = 10.5) -> None:
    p = d.add_paragraph()
    _run(p, text, size=size)


def note(d: Any, text: str) -> None:
    p = d.add_paragraph()
    _run(p, text, size=8, color=MUTED, italic=True)


def picture(d: Any, png: bytes, width_cm: float = 17.0) -> None:
    import io

    d.add_picture(io.BytesIO(png), width=Cm(width_cm))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


# ---------------------------------------------------------------------------------------------------- PowerPoint


def _prgb(hexcode: str) -> Any:
    from pptx.dml.color import RGBColor as PRGB

    return PRGB.from_string(hexcode)


def new_deck(template: Any = None) -> Any:
    """A 16:9 presentation; a client master is used as it is when supplied."""
    from pathlib import Path

    from pptx import Presentation
    from pptx.util import Inches

    if template and Path(template).exists():
        return Presentation(str(template))
    if _style().get("deck_bytes"):
        import io

        prs = Presentation(io.BytesIO(_style()["deck_bytes"]))
        ids = prs.slides._sldIdLst                     # the template's sample slides go; its master and layouts stay
        for sid in list(ids):
            prs.part.drop_rel(sid.rId)
            ids.remove(sid)
        return prs
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def _blank(prs: Any) -> Any:
    layouts = list(prs.slide_layouts)
    layout = next((lay for lay in layouts if lay.name.lower() == "blank"),
                  min(layouts, key=lambda lay: len(lay.placeholders)))     # a template's emptiest layout
    return prs.slides.add_slide(layout)


def _text(slide: Any, x: Any, y: Any, w: Any, h: Any, text: str, *, size: float = 14, bold: bool = False,
          color: str = TEXT, align: Any = None) -> Any:
    from pptx.util import Pt as PPt

    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.name = PPt(size), bold, font_name()
    r.font.color.rgb = _prgb(color)
    return box


def _rect(slide: Any, x: Any, y: Any, w: Any, h: Any, fill: str) -> Any:
    from pptx.enum.shapes import MSO_SHAPE

    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = _prgb(fill)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def title_slide(prs: Any, title: str, subtitle: str, meta: str) -> Any:
    from pptx.util import Inches

    s = _blank(prs)
    W, H = prs.slide_width, prs.slide_height
    _rect(s, 0, 0, W, H, primary())
    _rect(s, Inches(0.8), Inches(3.55), Inches(1.6), Inches(0.08), accent())
    _text(s, Inches(0.8), Inches(2.0), Inches(11.5), Inches(1.4), title, size=40, bold=True, color="FFFFFF")
    _text(s, Inches(0.8), Inches(3.8), Inches(11.5), Inches(0.6), subtitle, size=20, color="DCE3F0")
    _text(s, Inches(0.8), Inches(6.6), Inches(11.5), Inches(0.4), meta, size=11, color="AEB9CF")
    if logo():
        import io

        s.shapes.add_picture(io.BytesIO(logo()), Inches(0.8), Inches(0.6), height=Inches(0.8))
    return s


def content_slide(prs: Any, title: str, report: str, number: int, subtitle: str | None = None) -> Any:
    """A white slide with a title, an accent rule, a subtitle line and a footer (report name, classification,
    slide number)."""
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches

    s = _blank(prs)
    W = prs.slide_width
    _text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.7), title, size=28, bold=True, color=primary())
    _rect(s, Inches(0.6), Inches(1.05), Inches(1.2), Inches(0.06), accent())
    if subtitle:
        _text(s, Inches(0.6), Inches(1.18), Inches(12), Inches(0.45), subtitle, size=13, color=MUTED)
    _rect(s, 0, Inches(7.1), W, Inches(0.4), "F3F4F6")
    _text(s, Inches(0.6), Inches(7.16), Inches(9), Inches(0.3), f"{brand()}  ·  {report}  ·  {classification()}",
          size=9, color=MUTED)
    _text(s, Inches(11.7), Inches(7.16), Inches(1.0), Inches(0.3), str(number), size=9, color=MUTED,
          align=PP_ALIGN.RIGHT)
    return s


def deck_tiles(slide: Any, items: list[tuple[str, Any, str]], x: Any, y: Any, w: Any, *, h: Any = None,
               per_row: int = 3, unit_map: dict[str, str] | None = None) -> None:
    from pptx.util import Inches

    h = h or Inches(1.25)
    gap = Inches(0.2)
    tw = int((w - gap * (per_row - 1)) / per_row)
    for i, (label, value, tone) in enumerate(items):
        cx, cy = x + (i % per_row) * (tw + gap), y + (i // per_row) * (h + gap)
        _rect(slide, cx, cy, tw, h, TILE)
        _rect(slide, cx, cy, Inches(0.07), h, tone_colours().get(tone, primary()))
        _text(slide, cx + Inches(0.25), cy + Inches(0.12), tw - Inches(0.35), Inches(0.65),
              fmt(value, (unit_map or {}).get(label, ""), missing="–"), size=28, bold=True, color=tone_colours().get(tone, primary()))
        _text(slide, cx + Inches(0.25), cy + Inches(0.8), tw - Inches(0.35), Inches(0.4), label, size=11, color=MUTED)


def deck_table(slide: Any, header: list[str], rows: list[list[Any]], x: Any, y: Any, w: Any, *,
               col_widths: list[float] | None = None, tone_col: int | None = None, tones: dict[str, str] | None = None,
               numeric: set[int] | None = None, font: float = 11, empty: str = "Nothing to report.") -> None:
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches
    from pptx.util import Pt as PPt

    if not rows:
        _text(slide, x, y, w, Inches(0.4), empty, size=12, color=MUTED)
        return
    numeric = numeric or set()
    row_h = Inches(0.36)
    shape = slide.shapes.add_table(len(rows) + 1, len(header), x, y, w, row_h * (len(rows) + 1))
    tbl = shape.table
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = int(w * cw / total)
    for r_i, row in enumerate([header, *rows]):
        for c_i, v in enumerate(row):
            cell = tbl.cell(r_i, c_i)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _prgb(primary() if r_i == 0 else (ROW_ALT if r_i % 2 == 0 else "FFFFFF"))
            cell.margin_left = cell.margin_right = Inches(0.08)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if c_i in numeric else PP_ALIGN.LEFT
            run = p.add_run()
            run.text = str(v) if r_i == 0 else fmt(v, missing="–")
            run.font.size, run.font.name = PPt(font), font_name()
            color, bold = ("FFFFFF", True) if r_i == 0 else (TEXT, False)
            if r_i and tone_col is not None and c_i == tone_col:
                color, bold = tone_colours().get((tones or SEVERITY_TONE).get(str(v), "neutral"), TEXT), True
            run.font.color.rgb, run.font.bold = _prgb(color), bold


def deck_chart(slide: Any, chart: dict, x: Any, y: Any, w: Any, h: Any, *, colors: list[str] | None = None) -> Any:
    """A native bar chart in the house palette: categories top-down, values at the bars, whole-number axis, no
    chart junk. Editable in PowerPoint (the data is in the embedded sheet)."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Pt as PPt

    data = CategoryChartData()
    data.categories = chart["categories"]
    for srs in chart["series"]:
        data.add_series(srs["name"], srs["values"])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, x, y, w, h, data)
    ch = gf.chart
    ch.font.size, ch.font.name = PPt(11), font_name()
    ch.has_title = bool(chart.get("title"))
    if ch.has_title:
        ch.chart_title.text_frame.text = chart["title"]
        tp = ch.chart_title.text_frame.paragraphs[0]
        tp.runs[0].font.size, tp.runs[0].font.bold = PPt(13), True
        tp.runs[0].font.color.rgb = _prgb(primary())
    ch.has_legend = len(chart["series"]) > 1
    if ch.has_legend:
        ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    plot.data_labels.font.size = PPt(11)
    plot.data_labels.number_format, plot.data_labels.number_format_is_linked = "0", False
    ch.category_axis.reverse_order = True
    ch.category_axis.format.line.color.rgb = _prgb(BORDER)
    va = ch.value_axis
    peak = max((v for srs in chart["series"] for v in srs["values"]), default=1) or 1
    from soc_platform.reporting.charts import nice_step

    va.minimum_scale, va.major_unit = 0, nice_step(peak)              # whole numbers, steps a reader expects
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = _prgb("E5E7EB")
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = "0", False
    va.format.line.fill.background()
    from soc_platform.reporting.charts import STATUS

    if len(chart["series"]) == 1 and all(c in STATUS for c in chart["categories"]):
        _chart_colors(ch, list(chart["categories"]), per_point=True)    # severity / verdict bars in their colours
        return gf
    palette = colors or [accent(), "E08A1E", "3A9E6E", "B84A62", "6B5BA8"]
    for i, srs in enumerate(plot.series):
        srs.format.fill.solid()
        srs.format.fill.fore_color.rgb = _prgb(palette[i % len(palette)])
    return gf


def deck_note(slide: Any, text: str, x: Any, y: Any, w: Any, h: Any, *, size: float = 13) -> None:
    _text(slide, x, y, w, h, text, size=size, color=TEXT)


def picture_row(d: Any, pngs: list[bytes | None], *, total_width_cm: float = 17.4) -> None:
    """Charts side by side (a borderless table, one chart per cell); charts with no data are left out."""
    import io

    items = [p for p in pngs if p]
    if not items:
        return
    if len(items) == 1:
        picture(d, items[0], width_cm=min(total_width_cm, 15.0))
        return
    t = d.add_table(rows=1, cols=len(items))
    _no_borders(t)
    _cell_margins(t, 0, 0, 40, 40)
    w = total_width_cm / len(items) - 0.3
    for cell, png in zip(t.rows[0].cells, items, strict=True):
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run().add_picture(io.BytesIO(png), width=Cm(w))
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def _chart_colors(chart: Any, labels: list[str], per_point: bool) -> None:
    from soc_platform.reporting.charts import color_for

    plot = chart.plots[0]
    if per_point:
        for i, (pt_label, point) in enumerate(zip(labels, plot.series[0].points, strict=False)):
            r, g, b = color_for(pt_label, i)
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = _prgb(f"{r:02X}{g:02X}{b:02X}")
    else:
        for i, (name, srs) in enumerate(zip(labels, plot.series, strict=False)):
            r, g, b = color_for(name, i)
            srs.format.fill.solid()
            srs.format.fill.fore_color.rgb = _prgb(f"{r:02X}{g:02X}{b:02X}")


def deck_doughnut(slide: Any, title: str, counts: dict[str, Any], x: Any, y: Any, w: Any, h: Any, *,
                  order: tuple[str, ...] = ()) -> Any:
    """A native doughnut (share of each category), coloured by meaning (severity, verdict, priority)."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Pt as PPt

    keys = [k for k in order if counts.get(k)] + sorted(k for k in counts if k not in order and counts.get(k))
    if not keys:
        _text(slide, x, y, w, h, f"{title}: nothing in this period.", size=13, color=MUTED)
        return None
    data = CategoryChartData()
    data.categories = keys
    data.add_series(title, [counts[k] for k in keys])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, x, y, w, h, data)
    ch = gf.chart
    ch.font.size, ch.font.name = PPt(12), font_name()
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    tp = ch.chart_title.text_frame.paragraphs[0].runs[0]
    tp.font.size, tp.font.bold, tp.font.color.rgb = PPt(14), True, _prgb(primary())
    ch.has_legend = True
    ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.RIGHT, False
    plot = ch.plots[0]
    plot.has_data_labels = True
    plot.data_labels.show_value, plot.data_labels.show_category_name = True, False
    plot.data_labels.font.size, plot.data_labels.font.bold = PPt(13), True
    plot.data_labels.font.color.rgb = _prgb("FFFFFF")
    plot.data_labels.number_format, plot.data_labels.number_format_is_linked = "0", False
    _chart_colors(ch, keys, per_point=True)
    return gf


def deck_stacked(slide: Any, title: str, categories: list[str], series: list[tuple[str, list[float]]],
                 x: Any, y: Any, w: Any, h: Any) -> Any:
    """Native stacked horizontal bars (for example each team's findings within / past SLA)."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Pt as PPt

    if not categories:
        _text(slide, x, y, w, h, f"{title}: nothing in this period.", size=13, color=MUTED)
        return None
    data = CategoryChartData()
    data.categories = categories
    for name, vals in series:
        data.add_series(name, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, x, y, w, h, data)
    ch = gf.chart
    ch.font.size, ch.font.name = PPt(12), font_name()
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    tp = ch.chart_title.text_frame.paragraphs[0].runs[0]
    tp.font.size, tp.font.bold, tp.font.color.rgb = PPt(14), True, _prgb(primary())
    ch.has_legend = True
    ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    plot = ch.plots[0]
    plot.gap_width, plot.overlap = 50, 100
    plot.has_data_labels = True
    plot.data_labels.number_format, plot.data_labels.number_format_is_linked = '0;;;', False
    plot.data_labels.font.size, plot.data_labels.font.color.rgb = PPt(11), _prgb("FFFFFF")
    ch.category_axis.reverse_order = True
    peak = max((sum(v[i] for _, v in series) for i in range(len(categories))), default=1) or 1
    va = ch.value_axis
    from soc_platform.reporting.charts import nice_step

    va.minimum_scale, va.major_unit = 0, nice_step(peak)
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = "0", False
    va.major_gridlines.format.line.color.rgb = _prgb("E5E7EB")
    va.format.line.fill.background()
    _chart_colors(ch, [n for n, _ in series], per_point=False)
    return gf
